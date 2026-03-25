"""
FINAL BUILD — all application materials for JPMorgan Chase Transactions Specialist II.
Resume format: gray-blue sidebar, dark sidebar text, gold underlines, gray subsection bars.
Cover letter: JD-mirrored language.
"""
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.colors import HexColor
from reportlab.pdfgen import canvas
from reportlab.lib.utils import simpleSplit
import os

OUT = "/home/user/claude"
PAGE_W, PAGE_H = letter
SW = 170  # sidebar width in points (~2.36 inches)

# Colors from the template image
SBG = HexColor("#8B9CAD")   # gray-blue sidebar
SH  = HexColor("#1A1A1A")   # dark sidebar headers
ST  = HexColor("#1A1A1A")   # dark sidebar body text
GL  = HexColor("#C2A24D")   # gold accent line
GB  = HexColor("#DADADA")   # gray subsection bar
DK  = HexColor("#1A1A1A")   # main text
GY  = HexColor("#555555")   # gray caption


# ═══════════════════════════════════════════════════════════════
# RESUME PDF
# ═══════════════════════════════════════════════════════════════

def make_resume():
    fp = os.path.join(OUT, "Chase_Kinslow_Resume_JPMorganChase.pdf")
    c = canvas.Canvas(fp, pagesize=letter)

    # Sidebar background
    c.setFillColor(SBG)
    c.rect(0, 0, SW, PAGE_H, fill=1, stroke=0)

    # ──── SIDEBAR ────
    sy = PAGE_H - 36
    lm = 12  # left margin inside sidebar
    smw = SW - 24  # sidebar max text width

    # CONTACT
    c.setFillColor(SH); c.setFont("Helvetica-Bold", 7)
    c.drawString(lm, sy, "C O N T A C T"); sy -= 14
    c.setFillColor(ST); c.setFont("Helvetica", 7.5)
    c.drawString(lm, sy, "(501) 707-7779"); sy -= 10
    c.drawString(lm, sy, "chasekn30@yahoo.com"); sy -= 10
    c.drawString(lm, sy, "Monroe, LA 71201"); sy -= 14

    # CORE COMPETENCIES
    c.setFillColor(SH); c.setFont("Helvetica-Bold", 7)
    c.drawString(lm, sy, "C O R E   C O M P E T E N C I E S"); sy -= 12
    c.setFillColor(ST); c.setFont("Helvetica", 7)
    comps = [
        "Transaction Processing & Clearing",
        "Document Transaction Handling",
        "Data Extraction & System Input",
        "Production & Accuracy Standards",
        "Customer Inquiry Resolution",
        "Non-Routine Issue Escalation",
        "AI & Automation Proficiency",
        "Market Products & Regulations",
        "Compliance & Regulatory Adherence",
    ]
    for comp in comps:
        lines = simpleSplit(comp, "Helvetica", 7, smw - 12)
        for i, ln in enumerate(lines):
            if i == 0:
                c.drawString(lm, sy, "\u00BB" + ln)
            else:
                c.drawString(lm + 8, sy, ln)
            sy -= 10
    sy -= 8

    # CREDENTIALS
    c.setFillColor(SH); c.setFont("Helvetica-Bold", 7)
    c.drawString(lm, sy, "C R E D E N T I A L S"); sy -= 12
    c.setFillColor(ST); c.setFont("Helvetica-Bold", 7.5)
    c.drawString(lm, sy, "Licensed CPA"); sy -= 10
    c.setFont("Helvetica", 7.5)
    c.drawString(lm, sy, "Arkansas & Louisiana"); sy -= 12
    c.setFont("Helvetica-Bold", 7.5)
    c.drawString(lm, sy, "Arkansas Bar License"); sy -= 14

    # EDUCATION
    c.setFillColor(SH); c.setFont("Helvetica-Bold", 7)
    c.drawString(lm, sy, "E D U C A T I O N"); sy -= 12
    c.setFillColor(ST); c.setFont("Helvetica-Bold", 7.5)
    c.drawString(lm, sy, "Juris Doctor (J.D.)"); sy -= 10
    c.setFont("Helvetica", 7.5)
    c.drawString(lm, sy, "University of Arkansas"); sy -= 12
    c.setFont("Helvetica-Bold", 7.5)
    c.drawString(lm, sy, "Bachelor of Accountancy"); sy -= 10
    c.setFont("Helvetica", 7.5)
    c.drawString(lm, sy, "University of Mississippi"); sy -= 14

    # TOOLS
    c.setFillColor(SH); c.setFont("Helvetica-Bold", 7)
    c.drawString(lm, sy, "T O O L S"); sy -= 12
    c.setFillColor(ST); c.setFont("Helvetica", 7)
    tools = [
        "Advanced Excel & Financial Modeling",
        "QuickBooks Desktop & Online",
        "NCR POS Systems",
        "Digital Inventory Tracking",
        "Microsoft Office Suite",
        "ADP Workforce Now",
    ]
    for t in tools:
        lines = simpleSplit(t, "Helvetica", 7, smw - 12)
        for i, ln in enumerate(lines):
            if i == 0:
                c.drawString(lm, sy, "\u00BB" + ln)
            else:
                c.drawString(lm + 8, sy, ln)
            sy -= 10
    sy -= 8

    # ADDITIONAL
    c.setFillColor(SH); c.setFont("Helvetica-Bold", 7)
    c.drawString(lm, sy, "A D D I T I O N A L"); sy -= 12
    c.setFillColor(ST); c.setFont("Helvetica", 7)
    adds = [
        "GAAP, HUD & LIHTC Compliance",
        "Valid Driver\u2019s License \u2014 Clean MVR",
        "Available for Extended Hours",
    ]
    for a in adds:
        lines = simpleSplit(a, "Helvetica", 7, smw - 12)
        for i, ln in enumerate(lines):
            if i == 0:
                c.drawString(lm, sy, "\u00BB" + ln)
            else:
                c.drawString(lm + 8, sy, ln)
            sy -= 10

    # ──── MAIN CONTENT ────
    mx = SW + 18
    mw = PAGE_W - mx - 24
    my = PAGE_H - 36

    # Name
    c.setFont("Helvetica-Bold", 22); c.setFillColor(DK)
    c.drawString(mx, my, "CHASE KINSLOW"); my -= 14

    # Subtitle
    c.setFont("Helvetica", 8); c.setFillColor(GY)
    c.drawString(mx, my, "Compliance & Transaction Processing Professional"); my -= 18

    # ── PROFESSIONAL SUMMARY ──
    c.setFont("Helvetica-Bold", 10); c.setFillColor(DK)
    c.drawString(mx, my, "PROFESSIONAL SUMMARY"); my -= 4
    c.setStrokeColor(GL); c.setLineWidth(1.2)
    c.line(mx, my, mx + mw, my); my -= 12

    summary = (
        "CPA-licensed transaction processing and compliance professional with a Juris Doctor and "
        "Bachelor of Accountancy. Three years processing, clearing, and servicing document "
        "transactions involving moderately complex tasks \u2014 including regulatory submissions, "
        "financial reconciliations, and audit fieldwork where a single error disqualified an entire "
        "file. Maintains the highest standards of production and accuracy across high-volume "
        "transaction environments. Experienced handling customer inquiries, escalating non-routine "
        "issues to senior team members, and applying knowledge of industry regulations to ensure "
        "compliance in all transaction activities. Developing proficiency in AI and automation tools "
        "to optimize processes."
    )
    c.setFont("Helvetica", 8); c.setFillColor(DK)
    for ln in simpleSplit(summary, "Helvetica", 8, mw):
        c.drawString(mx, my, ln); my -= 11
    my -= 6

    # ── KEY ACHIEVEMENTS & SKILLS ──
    c.setFont("Helvetica-Bold", 10); c.setFillColor(DK)
    c.drawString(mx, my, "KEY ACHIEVEMENTS & SKILLS"); my -= 4
    c.setStrokeColor(GL); c.line(mx, my, mx + mw, my); my -= 12

    def subsec(title):
        nonlocal my
        c.setFillColor(GB)
        c.rect(mx - 2, my - 2, mw + 4, 12, fill=1, stroke=0)
        c.setFont("Helvetica-Bold", 8); c.setFillColor(DK)
        c.drawString(mx, my, title); my -= 14

    def bullet(text):
        nonlocal my
        c.setFont("Helvetica", 7.5); c.setFillColor(DK)
        c.drawString(mx + 2, my, "\u2022")
        lines = simpleSplit(text, "Helvetica", 7.5, mw - 14)
        for i, ln in enumerate(lines):
            if my < 40:
                c.showPage()
                c.setFillColor(SBG); c.rect(0, 0, SW, PAGE_H, fill=1, stroke=0)
                my = PAGE_H - 36
            c.setFont("Helvetica", 7.5); c.setFillColor(DK)
            c.drawString(mx + 12 if i == 0 else mx + 16, my, ln)
            my -= 10
        my -= 2

    # Section 1
    subsec("Document Transactions, Compliance & Accuracy")
    bullet("Processed and cleared document transactions involving moderately complex regulatory submissions, adhering to established routines and procedures across LIHTC, HUD, and GAAP frameworks while maintaining zero-error production and accuracy standards")
    bullet("Built parallel reconciliation models to extract, input, and track comprehensive financial data sets \u2014 reconciling cumulative construction draws and recommending variance-correcting adjustments to ensure total data integrity across client portfolios")
    bullet("Maintained up-to-date knowledge of market products and industry regulations (DSCR, FMR tiers, set-aside requirements), applying this knowledge to ensure compliance in all transaction and submission activities")
    my -= 4

    # Section 2
    subsec("Transaction Processing & Production Standards")
    bullet("Process and clear high-volume transactions daily, adhering to established routines and standard operating procedures while ensuring accuracy and timeliness across a fast-paced retail operation")
    bullet("Reduced cash-handling discrepancies by 15% through disciplined daily reconciliation of register transactions, applying the same servicing, researching, and settling discipline the role requires")
    bullet("Maintained inventory control systems by extracting data, inputting into digital tracking systems, and ensuring accurate record-keeping \u2014 reducing stockouts by 10% across all product categories")
    my -= 4

    # Section 3
    subsec("Customer Inquiries, Escalation & Process Optimization")
    bullet("Handled customer inquiries and requests across departments, providing a positive customer experience at all touchpoints while resolving service issues and improving resolution time by 10%")
    bullet("Escalated non-routine issues to senior team members, applying common sense and experience of similar situations to identify potential solutions \u2014 reducing service disruptions by 20% through proactive reporting")
    bullet("Contributed to process optimization by developing automated Excel tracking models and digital inventory systems that streamlined daily operations and reduced manual data entry errors")
    my -= 8

    # ── EMPLOYMENT HISTORY ──
    c.setFont("Helvetica-Bold", 10); c.setFillColor(DK)
    c.drawString(mx, my, "EMPLOYMENT HISTORY"); my -= 4
    c.setStrokeColor(GL); c.line(mx, my, mx + mw, my); my -= 14

    for title, company, dates in [
        ("Operations & Logistics", "Bayou DeSiard Country Club", "2025"),
        ("Accountant", "Little and Associates, LLC", "2020 \u2013 2023"),
        ("Sales Associate", "Office Depot", "2020 \u2013 Present"),
    ]:
        c.setFont("Helvetica-Bold", 8); c.setFillColor(DK)
        c.drawString(mx, my, title)
        tw = c.stringWidth(title, "Helvetica-Bold", 8)
        c.setFont("Helvetica-Oblique", 8); c.setFillColor(GY)
        c.drawString(mx + tw, my, f" | {company}")
        c.setFont("Helvetica", 8)
        dw = c.stringWidth(dates, "Helvetica", 8)
        c.drawString(mx + mw - dw, my, dates)
        my -= 12

    c.save()
    print(f"Resume PDF: {os.path.getsize(fp):,} bytes")


# ═══════════════════════════════════════════════════════════════
# COVER LETTER PDF — rewritten to mirror JD language precisely
# ═══════════════════════════════════════════════════════════════

def make_cover_letter():
    fp = os.path.join(OUT, "Chase_Kinslow_CoverLetter_JPMorganChase.pdf")
    c = canvas.Canvas(fp, pagesize=letter)
    x, y = 72, PAGE_H - 72
    mw = PAGE_W - 144

    def wl(text, font="Helvetica", size=10.5, sp=14, color=DK):
        nonlocal y
        c.setFont(font, size); c.setFillColor(color)
        if not text.strip(): y -= sp; return
        for ln in simpleSplit(text, font, size, mw):
            c.drawString(x, y, ln); y -= sp

    wl("Chase Kinslow", "Helvetica-Bold", 14, 16)
    wl("Monroe, LA 71201", "Helvetica", 9.5, 13, GY)
    wl("(501) 707-7779  |  chasekn30@yahoo.com", "Helvetica", 9.5, 13, GY)
    y -= 16
    wl("March 25, 2026"); y -= 8
    wl("JPMorgan Chase & Co.")
    wl("Hiring Manager \u2014 Transactions Specialist II")
    wl("Monroe, LA"); y -= 8
    wl("Dear Hiring Manager,"); y -= 4

    paras = [
        # Para 1: Mirror JD opening — "processing and servicing transactions, handling document transactions, moderately complex tasks, structured and supervised environment"
        "I am writing to apply for the Transactions Specialist II position at the Monroe Operations Center. With a CPA license, a Juris Doctor, and a Bachelor of Accountancy, I have three years of direct experience processing, clearing, and servicing document transactions involving moderately complex tasks \u2014 and I am ready to contribute to your team within the structured and supervised environment this role operates in.",

        # Para 2: Mirror JD — "extracting checks and remittances, inputting data into systems, highest standards of production and accuracy"
        "I currently process and clear high-volume transactions daily at Office Depot, adhering to established routines and standard operating procedures while maintaining the highest standards of production and accuracy. I manage daily register reconciliation, cash handling, and inventory control by extracting data and inputting it into digital tracking systems \u2014 reducing cash discrepancies by 15% and stockouts by 10%. This is the same discipline required for extracting checks and remittances from envelopes, inputting data into systems, and maintaining accuracy under production pressure.",

        # Para 3: Mirror JD — "market products and industry regulations, compliance in all transaction activities"
        "Previously, as an Accountant at Little and Associates, LLC, I processed and cleared document transactions across LIHTC, HUD, and GAAP regulatory frameworks where a single discrepancy disqualified an entire project file. I built parallel reconciliation models to track financial data sets, managed comprehensive audit fieldwork including GAAP-compliant adjustments and HUD-tailored supplemental schedules, and maintained up-to-date knowledge of market products and industry regulations \u2014 applying that knowledge to ensure compliance in all transaction and submission activities.",

        # Para 4: Mirror JD — "customer inquiries, positive customer experience, escalate non-routine issues, common sense and experience, AI and automation"
        "I am experienced handling customer inquiries and requests, providing a positive customer experience at all touchpoints. I escalate non-routine issues to senior team members by applying common sense and experience of similar situations to identify potential solutions \u2014 an approach that reduced service disruptions by 20% in my most recent operations role. I am also developing proficiency in AI and automation tools to optimize transaction processes, and I am eager to contribute to JPMorgan Chase\u2019s innovation efforts as the firm scales toward 1,000+ AI use cases through tools like the LLM Suite.",

        # Para 5: Close — local, physical requirements, Monroe-specific knowledge
        "I understand the Monroe campus serves as a hub for global document custody and collateral review across home loans, auto loans, and securities, with the new $31 million Ruston Operations Center expanding that capacity. I am local, available immediately, and comfortable with the physical requirements of the role. I would welcome the chance to discuss how my background fits your team.",

        "Thank you for your consideration.",
    ]
    for p in paras:
        wl(p); y -= 6

    y -= 4
    wl("Sincerely,")
    wl("Chase Kinslow", "Helvetica-Bold", 10.5)

    c.save()
    print(f"Cover Letter PDF: {os.path.getsize(fp):,} bytes")


# ═══════════════════════════════════════════════════════════════
# COVER LETTER DOCX
# ═══════════════════════════════════════════════════════════════

def make_cover_letter_docx():
    from docx import Document
    from docx.shared import Pt, Inches

    COVER_TEXT = """Chase Kinslow
Monroe, LA 71201
(501) 707-7779 | chasekn30@yahoo.com

March 25, 2026

JPMorgan Chase & Co.
Hiring Manager \u2014 Transactions Specialist II
Monroe, LA

Dear Hiring Manager,

I am writing to apply for the Transactions Specialist II position at the Monroe Operations Center. With a CPA license, a Juris Doctor, and a Bachelor of Accountancy, I have three years of direct experience processing, clearing, and servicing document transactions involving moderately complex tasks \u2014 and I am ready to contribute to your team within the structured and supervised environment this role operates in.

I currently process and clear high-volume transactions daily at Office Depot, adhering to established routines and standard operating procedures while maintaining the highest standards of production and accuracy. I manage daily register reconciliation, cash handling, and inventory control by extracting data and inputting it into digital tracking systems \u2014 reducing cash discrepancies by 15% and stockouts by 10%. This is the same discipline required for extracting checks and remittances from envelopes, inputting data into systems, and maintaining accuracy under production pressure.

Previously, as an Accountant at Little and Associates, LLC, I processed and cleared document transactions across LIHTC, HUD, and GAAP regulatory frameworks where a single discrepancy disqualified an entire project file. I built parallel reconciliation models to track financial data sets, managed comprehensive audit fieldwork including GAAP-compliant adjustments and HUD-tailored supplemental schedules, and maintained up-to-date knowledge of market products and industry regulations \u2014 applying that knowledge to ensure compliance in all transaction and submission activities.

I am experienced handling customer inquiries and requests, providing a positive customer experience at all touchpoints. I escalate non-routine issues to senior team members by applying common sense and experience of similar situations to identify potential solutions \u2014 an approach that reduced service disruptions by 20% in my most recent operations role. I am also developing proficiency in AI and automation tools to optimize transaction processes, and I am eager to contribute to JPMorgan Chase\u2019s innovation efforts as the firm scales toward 1,000+ AI use cases through tools like the LLM Suite.

I understand the Monroe campus serves as a hub for global document custody and collateral review across home loans, auto loans, and securities, with the new $31 million Ruston Operations Center expanding that capacity. I am local, available immediately, and comfortable with the physical requirements of the role. I would welcome the chance to discuss how my background fits your team.

Thank you for your consideration.

Sincerely,
Chase Kinslow"""

    doc = Document()
    style = doc.styles['Normal']; style.font.name = 'Calibri'; style.font.size = Pt(10.5)
    for s in doc.sections:
        s.top_margin = Inches(1); s.bottom_margin = Inches(1)
        s.left_margin = Inches(1); s.right_margin = Inches(1)

    for line in COVER_TEXT.strip().split('\n'):
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(2)
        if line == "Chase Kinslow":
            if p.runs: p.runs[0].bold = True; p.runs[0].font.size = Pt(14)

    fp = os.path.join(OUT, "Chase_Kinslow_CoverLetter_JPMorganChase.docx")
    doc.save(fp)
    print(f"Cover Letter DOCX: {os.path.getsize(fp):,} bytes")


# ═══════════════════════════════════════════════════════════════
# RESUME DOCX — two-column table layout
# ═══════════════════════════════════════════════════════════════

def make_resume_docx():
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    style = doc.styles['Normal']; style.font.name = 'Calibri'; style.font.size = Pt(9)
    for s in doc.sections:
        s.top_margin = Inches(0.3); s.bottom_margin = Inches(0.3)
        s.left_margin = Inches(0.3); s.right_margin = Inches(0.3)

    # Create a 1-row, 2-column table
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.4)
    table.columns[1].width = Inches(5.1)

    # Remove table borders
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'none'); el.set(qn('w:sz'), '0'); el.set(qn('w:space'), '0'); el.set(qn('w:color'), 'auto')
        borders.append(el)
    tblPr.append(borders)

    left = table.cell(0, 0)
    right = table.cell(0, 1)

    # Shade left cell gray-blue
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear'); shading.set(qn('w:color'), 'auto'); shading.set(qn('w:fill'), '8B9CAD')
    left._element.get_or_add_tcPr().append(shading)

    SBC = RGBColor(0x1A, 0x1A, 0x1A)
    D = RGBColor(0x1A, 0x1A, 0x1A)
    G = RGBColor(0x55, 0x55, 0x55)

    # ── LEFT COLUMN (sidebar) ──
    def lsec(cell, title):
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(4)
        r = p.add_run(title); r.font.size = Pt(7); r.bold = True; r.font.color.rgb = SBC

    def litm(cell, text, bold=False):
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(text)
        r.font.size = Pt(7.5); r.font.color.rgb = SBC
        if bold: r.bold = True

    def larr(cell, text):
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run("\u00BB " + text); r.font.size = Pt(7); r.font.color.rgb = SBC

    # Clear default empty paragraph
    left.paragraphs[0].text = ""

    lsec(left, "C O N T A C T")
    litm(left, "(501) 707-7779")
    litm(left, "chasekn30@yahoo.com")
    litm(left, "Monroe, LA 71201")

    lsec(left, "C O R E   C O M P E T E N C I E S")
    for c2 in ["Transaction Processing & Clearing", "Document Transaction Handling", "Data Extraction & System Input",
               "Production & Accuracy Standards", "Customer Inquiry Resolution", "Non-Routine Issue Escalation",
               "AI & Automation Proficiency", "Market Products & Regulations", "Compliance & Regulatory Adherence"]:
        larr(left, c2)

    lsec(left, "C R E D E N T I A L S")
    litm(left, "Licensed CPA", bold=True)
    litm(left, "Arkansas & Louisiana")
    litm(left, "Arkansas Bar License", bold=True)

    lsec(left, "E D U C A T I O N")
    litm(left, "Juris Doctor (J.D.)", bold=True)
    litm(left, "University of Arkansas")
    litm(left, "Bachelor of Accountancy", bold=True)
    litm(left, "University of Mississippi")

    lsec(left, "T O O L S")
    for t in ["Advanced Excel & Financial Modeling", "QuickBooks Desktop & Online", "NCR POS Systems",
              "Digital Inventory Tracking", "Microsoft Office Suite", "ADP Workforce Now"]:
        larr(left, t)

    lsec(left, "A D D I T I O N A L")
    for a in ["GAAP, HUD & LIHTC Compliance", "Valid Driver\u2019s License \u2014 Clean MVR", "Available for Extended Hours"]:
        larr(left, a)

    # ── RIGHT COLUMN (main content) ──
    right.paragraphs[0].text = ""

    def rsec(cell, title):
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title); r.font.size = Pt(10); r.bold = True; r.font.color.rgb = D
        # Gold bottom border
        pPr = p._element.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        b = OxmlElement('w:bottom'); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '6')
        b.set(qn('w:space'), '1'); b.set(qn('w:color'), 'C2A24D')
        pBdr.append(b); pPr.append(pBdr)

    def rsubsec(cell, title):
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(2)
        pPr = p._element.get_or_add_pPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'DADADA')
        pPr.append(shd)
        r = p.add_run(title); r.font.size = Pt(8); r.bold = True; r.font.color.rgb = D

    def rbul(cell, text):
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.left_indent = Pt(12)
        r = p.add_run("\u2022 " + text); r.font.size = Pt(7.5); r.font.color.rgb = D

    # Name
    p = right.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("CHASE KINSLOW"); r.font.size = Pt(22); r.bold = True; r.font.color.rgb = D

    p = right.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Compliance & Transaction Processing Professional"); r.font.size = Pt(8); r.font.color.rgb = G

    rsec(right, "PROFESSIONAL SUMMARY")
    p = right.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(
        "CPA-licensed transaction processing and compliance professional with a Juris Doctor and "
        "Bachelor of Accountancy. Three years processing, clearing, and servicing document "
        "transactions involving moderately complex tasks \u2014 including regulatory submissions, "
        "financial reconciliations, and audit fieldwork where a single error disqualified an entire "
        "file. Maintains the highest standards of production and accuracy across high-volume "
        "transaction environments. Experienced handling customer inquiries, escalating non-routine "
        "issues to senior team members, and applying knowledge of industry regulations to ensure "
        "compliance in all transaction activities. Developing proficiency in AI and automation tools "
        "to optimize processes."
    ); r.font.size = Pt(8)

    rsec(right, "KEY ACHIEVEMENTS & SKILLS")

    rsubsec(right, "Document Transactions, Compliance & Accuracy")
    for b in [
        "Processed and cleared document transactions involving moderately complex regulatory submissions, adhering to established routines and procedures across LIHTC, HUD, and GAAP frameworks while maintaining zero-error production and accuracy standards",
        "Built parallel reconciliation models to extract, input, and track comprehensive financial data sets \u2014 reconciling cumulative construction draws and recommending variance-correcting adjustments to ensure total data integrity across client portfolios",
        "Maintained up-to-date knowledge of market products and industry regulations (DSCR, FMR tiers, set-aside requirements), applying this knowledge to ensure compliance in all transaction and submission activities",
    ]: rbul(right, b)

    rsubsec(right, "Transaction Processing & Production Standards")
    for b in [
        "Process and clear high-volume transactions daily, adhering to established routines and standard operating procedures while ensuring accuracy and timeliness across a fast-paced retail operation",
        "Reduced cash-handling discrepancies by 15% through disciplined daily reconciliation of register transactions, applying the same servicing, researching, and settling discipline the role requires",
        "Maintained inventory control systems by extracting data, inputting into digital tracking systems, and ensuring accurate record-keeping \u2014 reducing stockouts by 10% across all product categories",
    ]: rbul(right, b)

    rsubsec(right, "Customer Inquiries, Escalation & Process Optimization")
    for b in [
        "Handled customer inquiries and requests across departments, providing a positive customer experience at all touchpoints while resolving service issues and improving resolution time by 10%",
        "Escalated non-routine issues to senior team members, applying common sense and experience of similar situations to identify potential solutions \u2014 reducing service disruptions by 20% through proactive reporting",
        "Contributed to process optimization by developing automated Excel tracking models and digital inventory systems that streamlined daily operations and reduced manual data entry errors",
    ]: rbul(right, b)

    rsec(right, "EMPLOYMENT HISTORY")
    for title, company, dates in [
        ("Operations & Logistics", "Bayou DeSiard Country Club", "2025"),
        ("Accountant", "Little and Associates, LLC", "2020 \u2013 2023"),
        ("Sales Associate", "Office Depot", "2020 \u2013 Present"),
    ]:
        p = right.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title); r.bold = True; r.font.size = Pt(8); r.font.color.rgb = D
        r = p.add_run(f" | {company}"); r.font.size = Pt(8); r.italic = True; r.font.color.rgb = G
        r = p.add_run(f"    {dates}"); r.font.size = Pt(8); r.font.color.rgb = G

    fp2 = os.path.join(OUT, "Chase_Kinslow_Resume_JPMorganChase.docx")
    doc.save(fp2)
    print(f"Resume DOCX: {os.path.getsize(fp2):,} bytes")


if __name__ == "__main__":
    make_resume()
    make_resume_docx()
    make_cover_letter()
    make_cover_letter_docx()
    print("\nAll files generated!")
