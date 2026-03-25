"""
JPMorgan Chase Transactions Specialist II Resume
FORMAT: Midsouth Medical template (dark navy sidebar, orange accents)
CONTENT: JD-tailored bullets mirroring exact JD language
"""
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.colors import HexColor, white
from reportlab.pdfgen import canvas
from reportlab.lib.utils import simpleSplit
import os

OUT = "/home/user/claude"

# Midsouth Medical template colors
NAVY = HexColor("#2B3544")
ORANGE = HexColor("#E8872B")
DARK = HexColor("#1A1A1A")
WHITE = white
PAGE_W, PAGE_H = letter
SW = 2.1 * inch  # sidebar width


def sbg(c):
    c.setFillColor(NAVY)
    c.rect(0, 0, SW, PAGE_H, fill=1, stroke=0)

def ssec(c, title, y):
    c.setFillColor(ORANGE); c.setFont("Helvetica-Bold", 9.5)
    c.drawString(18, y, title)
    return y - 16

def sitm(c, text, y, bold=False):
    c.setFillColor(WHITE)
    f = "Helvetica-Bold" if bold else "Helvetica"
    c.setFont(f, 8.5)
    for ln in simpleSplit(text, f, 8.5, SW - 36):
        c.drawString(22, y, ln); y -= 12
    return y

def sbul(c, text, y):
    c.setFillColor(WHITE); c.setFont("Helvetica", 8.5)
    lines = simpleSplit(text, "Helvetica", 8.5, SW - 44)
    for i, ln in enumerate(lines):
        if i == 0:
            c.drawString(22, y, "\u2022  " + ln)
        else:
            c.drawString(32, y, ln)
        y -= 12
    return y

def msec(c, title, y, mx, mw):
    c.setFont("Helvetica-Bold", 11.5); c.setFillColor(DARK)
    c.drawString(mx, y, title); y -= 4
    c.setStrokeColor(ORANGE); c.setLineWidth(1.5)
    c.line(mx, y, mx + mw, y)
    return y - 14

def msubsec(c, title, y, mx):
    c.setFont("Helvetica-Bold", 10); c.setFillColor(DARK)
    c.drawString(mx, y, title)
    return y - 14

def mbul(c, text, x, y, mw):
    c.setFillColor(ORANGE); c.setFont("Helvetica", 10)
    c.drawString(x, y, "\u2014")
    c.setFillColor(DARK); c.setFont("Helvetica", 9.5)
    tx = x + 12
    for i, ln in enumerate(simpleSplit(text, "Helvetica", 9.5, mw - 12)):
        if y < 40: c.showPage(); y = PAGE_H - 40; sbg(c)
        c.setFont("Helvetica", 9.5); c.setFillColor(DARK)
        c.drawString(tx if i == 0 else tx + 4, y, ln); y -= 13
    return y

def mwrap(c, text, x, y, mw):
    c.setFont("Helvetica", 9.5); c.setFillColor(DARK)
    for ln in simpleSplit(text, "Helvetica", 9.5, mw):
        if y < 40: c.showPage(); y = PAGE_H - 40; sbg(c)
        c.setFont("Helvetica", 9.5); c.setFillColor(DARK)
        c.drawString(x, y, ln); y -= 13
    return y


def make_resume_pdf():
    fp = os.path.join(OUT, "Chase_Kinslow_Resume_JPMorganChase.pdf")
    c = canvas.Canvas(fp, pagesize=letter)
    sbg(c)

    # ════════════ SIDEBAR ════════════
    sy = PAGE_H - 50

    sy = ssec(c, "CONTACT", sy)
    sy = sitm(c, "(501) 707-7779", sy)
    sy = sitm(c, "chasekn30@yahoo.com", sy)
    sy = sitm(c, "Monroe, LA 71201", sy)
    sy -= 10

    sy = ssec(c, "CORE COMPETENCIES", sy)
    for x in [
        "Transaction Processing & Clearing",
        "Document Transaction Handling",
        "Data Extraction & System Input",
        "Production & Accuracy Standards",
        "Customer Inquiry Resolution",
        "Non-Routine Issue Escalation",
        "AI & Automation Proficiency",
        "Market Products & Regulations",
        "Compliance & Regulatory Adherence",
    ]:
        sy = sbul(c, x, sy)
    sy -= 10

    sy = ssec(c, "CREDENTIALS", sy)
    sy = sitm(c, "Licensed CPA", sy, bold=True)
    sy = sitm(c, "Arkansas & Louisiana", sy)
    sy -= 4
    sy = sitm(c, "Arkansas Bar License", sy, bold=True)
    sy -= 10

    sy = ssec(c, "EDUCATION", sy)
    sy = sitm(c, "Juris Doctor (J.D.)", sy, bold=True)
    sy = sitm(c, "University of Arkansas", sy)
    sy -= 4
    sy = sitm(c, "Bachelor of Accountancy", sy, bold=True)
    sy = sitm(c, "University of Mississippi", sy)
    sy -= 10

    sy = ssec(c, "TOOLS", sy)
    for x in [
        "Advanced Excel & Financial Modeling",
        "QuickBooks Desktop & Online",
        "NCR POS Systems",
        "Digital Inventory Tracking",
        "Microsoft Office Suite",
        "ADP Workforce Now",
    ]:
        sy = sbul(c, x, sy)

    # ════════════ MAIN CONTENT ════════════
    mx = SW + 24
    mw = PAGE_W - mx - 30
    my = PAGE_H - 48

    c.setFont("Helvetica-Bold", 26); c.setFillColor(DARK)
    c.drawString(mx, my, "CHASE KINSLOW"); my -= 22

    c.setFont("Helvetica-Oblique", 11); c.setFillColor(ORANGE)
    c.drawString(mx, my, "Compliance & Transaction Processing Professional"); my -= 28

    # ── PROFESSIONAL SUMMARY ──
    my = msec(c, "PROFESSIONAL SUMMARY", my, mx, mw)
    my = mwrap(c, (
        "CPA-licensed transaction processing and compliance professional with a Juris Doctor and "
        "Bachelor of Accountancy. Three years processing, clearing, and servicing document "
        "transactions involving moderately complex tasks \u2014 including regulatory submissions, "
        "financial reconciliations, and audit fieldwork where a single error disqualified an entire "
        "file. Maintains the highest standards of production and accuracy across high-volume "
        "transaction environments. Experienced handling customer inquiries, escalating non-routine "
        "issues to senior team members, and applying knowledge of industry regulations to ensure "
        "compliance in all transaction activities. Developing proficiency in AI and automation tools "
        "to optimize processes."
    ), mx, my, mw); my -= 10

    # ── KEY ACHIEVEMENTS & SKILLS ──
    my = msec(c, "KEY ACHIEVEMENTS & SKILLS", my, mx, mw)

    my = msubsec(c, "Document Transactions, Compliance & Accuracy", my, mx)
    for b in [
        "Processed and cleared document transactions involving moderately complex regulatory submissions, adhering to established routines and procedures across LIHTC, HUD, and GAAP frameworks while maintaining zero-error production and accuracy standards",
        "Built parallel reconciliation models to extract, input, and track comprehensive financial data sets \u2014 reconciling cumulative construction draws and recommending variance-correcting adjustments to ensure total data integrity across client portfolios",
        "Maintained up-to-date knowledge of market products and industry regulations (DSCR, FMR tiers, set-aside requirements), applying this knowledge to ensure compliance in all transaction and submission activities",
    ]:
        my = mbul(c, b, mx, my, mw); my -= 2
    my -= 6

    my = msubsec(c, "Transaction Processing & Production Standards", my, mx)
    for b in [
        "Process and clear high-volume transactions daily, adhering to established routines and standard operating procedures while ensuring accuracy and timeliness across a fast-paced retail operation",
        "Reduced cash-handling discrepancies by 15% through disciplined daily reconciliation of register transactions, applying the same servicing, researching, and settling discipline the role requires",
        "Maintained inventory control systems by extracting data, inputting into digital tracking systems, and ensuring accurate record-keeping \u2014 reducing stockouts by 10% across all product categories",
    ]:
        my = mbul(c, b, mx, my, mw); my -= 2
    my -= 6

    my = msubsec(c, "Customer Inquiries, Escalation & Process Optimization", my, mx)
    for b in [
        "Handled customer inquiries and requests across departments, providing a positive customer experience at all touchpoints while resolving service issues and improving resolution time by 10%",
        "Escalated non-routine issues to senior team members, applying common sense and experience of similar situations to identify potential solutions \u2014 reducing service disruptions by 20% through proactive reporting",
        "Contributed to process optimization by developing automated Excel tracking models and digital inventory systems that streamlined daily operations and reduced manual data entry errors",
    ]:
        my = mbul(c, b, mx, my, mw); my -= 2
    my -= 10

    # ── EMPLOYMENT HISTORY ──
    my = msec(c, "EMPLOYMENT HISTORY", my, mx, mw)
    for title, company, dates in [
        ("Operations & Logistics", "Bayou DeSiard Country Club", "2025"),
        ("Accountant", "Little and Associates, LLC", "2020 \u2013 2023"),
        ("Sales Associate", "Office Depot", "2020 \u2013 Present"),
    ]:
        c.setFont("Helvetica-Bold", 9.5); c.setFillColor(DARK)
        lt = f"{title}  |  "; c.drawString(mx, my, lt)
        tw = c.stringWidth(lt, "Helvetica-Bold", 9.5)
        c.setFont("Helvetica-Oblique", 9.5); c.drawString(mx + tw, my, company)
        tw2 = c.stringWidth(company, "Helvetica-Oblique", 9.5)
        c.setFont("Helvetica", 9.5); c.drawString(mx + tw + tw2, my, f"  |  {dates}")
        my -= 16

    c.save()
    print(f"Resume PDF: {fp} ({os.path.getsize(fp):,} bytes)")


def make_resume_docx():
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    style = doc.styles['Normal']; style.font.name = 'Calibri'; style.font.size = Pt(10)
    for s in doc.sections:
        s.top_margin = Inches(0.5); s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.6); s.right_margin = Inches(0.6)

    O = RGBColor(0xE8, 0x87, 0x2B); D = RGBColor(0x1A, 0x1A, 0x1A); G = RGBColor(0x55, 0x55, 0x55)

    def orange_line(doc):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(4)
        pPr = p._element.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr')
        b = OxmlElement('w:bottom'); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '6')
        b.set(qn('w:space'), '1'); b.set(qn('w:color'), 'E8872B')
        pBdr.append(b); pPr.append(pBdr)

    def sec(doc, t):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t); r.font.size = Pt(12); r.bold = True; r.font.color.rgb = D; orange_line(doc)

    def subsec(doc, t):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(t); r.font.size = Pt(10.5); r.bold = True; r.font.color.rgb = D

    def bul(doc, t):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.left_indent = Pt(18)
        r = p.add_run("\u2014  " + t); r.font.size = Pt(9.5); r.font.color.rgb = D

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after = Pt(0)
    r = p.add_run("CHASE KINSLOW"); r.font.size = Pt(24); r.bold = True; r.font.color.rgb = D

    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Compliance & Transaction Processing Professional"); r.font.size = Pt(11); r.italic = True; r.font.color.rgb = O

    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
    r = p.add_run("(501) 707-7779  |  chasekn30@yahoo.com  |  Monroe, LA 71201"); r.font.size = Pt(9); r.font.color.rgb = G

    sec(doc, "PROFESSIONAL SUMMARY")
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(
        "CPA-licensed transaction processing and compliance professional with a Juris Doctor and "
        "Bachelor of Accountancy. Three years processing, clearing, and servicing document "
        "transactions involving moderately complex tasks \u2014 including regulatory submissions, "
        "financial reconciliations, and audit fieldwork where a single error disqualified an entire "
        "file. Maintains the highest standards of production and accuracy across high-volume "
        "transaction environments. Experienced handling customer inquiries, escalating non-routine "
        "issues to senior team members, and applying knowledge of industry regulations to ensure "
        "compliance in all transaction activities. Developing proficiency in AI and automation tools "
        "to optimize processes."
    ); r.font.size = Pt(9.5)

    sec(doc, "KEY ACHIEVEMENTS & SKILLS")

    subsec(doc, "Document Transactions, Compliance & Accuracy")
    for b in [
        "Processed and cleared document transactions involving moderately complex regulatory submissions, adhering to established routines and procedures across LIHTC, HUD, and GAAP frameworks while maintaining zero-error production and accuracy standards",
        "Built parallel reconciliation models to extract, input, and track comprehensive financial data sets \u2014 reconciling cumulative construction draws and recommending variance-correcting adjustments to ensure total data integrity across client portfolios",
        "Maintained up-to-date knowledge of market products and industry regulations (DSCR, FMR tiers, set-aside requirements), applying this knowledge to ensure compliance in all transaction and submission activities",
    ]: bul(doc, b)

    subsec(doc, "Transaction Processing & Production Standards")
    for b in [
        "Process and clear high-volume transactions daily, adhering to established routines and standard operating procedures while ensuring accuracy and timeliness across a fast-paced retail operation",
        "Reduced cash-handling discrepancies by 15% through disciplined daily reconciliation of register transactions, applying the same servicing, researching, and settling discipline the role requires",
        "Maintained inventory control systems by extracting data, inputting into digital tracking systems, and ensuring accurate record-keeping \u2014 reducing stockouts by 10% across all product categories",
    ]: bul(doc, b)

    subsec(doc, "Customer Inquiries, Escalation & Process Optimization")
    for b in [
        "Handled customer inquiries and requests across departments, providing a positive customer experience at all touchpoints while resolving service issues and improving resolution time by 10%",
        "Escalated non-routine issues to senior team members, applying common sense and experience of similar situations to identify potential solutions \u2014 reducing service disruptions by 20% through proactive reporting",
        "Contributed to process optimization by developing automated Excel tracking models and digital inventory systems that streamlined daily operations and reduced manual data entry errors",
    ]: bul(doc, b)

    sec(doc, "EMPLOYMENT HISTORY")
    for title, company, dates in [
        ("Operations & Logistics", "Bayou DeSiard Country Club", "2025"),
        ("Accountant", "Little and Associates, LLC", "2020 \u2013 2023"),
        ("Sales Associate", "Office Depot", "2020 \u2013 Present"),
    ]:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = D
        r = p.add_run("  |  "); r.font.size = Pt(9.5)
        r = p.add_run(company); r.italic = True; r.font.size = Pt(9.5)
        r = p.add_run(f"  |  {dates}"); r.font.size = Pt(9.5)

    sec(doc, "CREDENTIALS & EDUCATION")
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Licensed CPA (Arkansas & Louisiana)  \u2022  Arkansas Bar License  \u2022  J.D., University of Arkansas  \u2022  B.Accy., University of Mississippi")
    r.font.size = Pt(9)

    sec(doc, "TOOLS")
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Advanced Excel & Financial Modeling  \u2022  QuickBooks Desktop & Online  \u2022  NCR POS Systems  \u2022  Digital Inventory Tracking  \u2022  Microsoft Office Suite  \u2022  ADP Workforce Now")
    r.font.size = Pt(9)

    fp2 = os.path.join(OUT, "Chase_Kinslow_Resume_JPMorganChase.docx")
    doc.save(fp2)
    print(f"Resume DOCX: {fp2} ({os.path.getsize(fp2):,} bytes)")


if __name__ == "__main__":
    make_resume_pdf()
    make_resume_docx()
    print("Done!")
