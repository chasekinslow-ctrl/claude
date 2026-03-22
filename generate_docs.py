from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.colors import HexColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import os

OUT = "/home/user/claude"

# ─── Helper: Word ───────────────────────────────────────────────

def add_heading_styled(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    return h

def add_para(doc, text, bold=False, size=11, align=None, space_after=4):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p

def add_bullet(doc, text, size=11):
    p = doc.add_paragraph(style='List Bullet')
    p.clear()
    run = p.add_run(text)
    run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(2)

# ─── RESUME (Word) ──────────────────────────────────────────────

def make_resume_docx():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    add_para(doc, "CHASE KINSLOW", bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, "Monroe, LA  |  Valid Driver\u2019s License", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

    add_heading_styled(doc, "Professional Summary", level=2)
    add_para(doc, "Dedicated operations professional with hands-on experience in facility maintenance, sanitation, logistics, and customer service. Proven ability to maintain clean, organized workspaces while upholding safety and compliance standards. Strong attention to detail with a track record of reliability, punctuality, and team collaboration across fast-paced environments.")

    add_heading_styled(doc, "Work Experience", level=2)

    add_para(doc, "Operations & Logistics  |  Bayou DeSiard Country Club", bold=True, size=11, space_after=1)
    add_para(doc, "May 2025 \u2013 Present", size=10, space_after=4)
    for b in [
        "Managed daily facility operations including cleaning, sanitation, and maintenance of indoor and outdoor spaces to uphold premium presentation standards",
        "Performed equipment inspections and ensured all tools and machinery were maintained in safe, working condition",
        "Coordinated logistics for events and daily operations, demonstrating strong organizational and multitasking skills",
        "Maintained clean and orderly work areas, identifying and reporting unsafe or hazardous conditions",
        "Collaborated with team members to deliver exceptional service to club members and guests",
    ]:
        add_bullet(doc, b)

    add_para(doc, "Operations & Compliance Associate  |  Little and Associates", bold=True, size=11, space_after=1)
    add_para(doc, "December 2020 \u2013 2025", size=10, space_after=4)
    for b in [
        "Executed daily operational tasks with strict adherence to regulatory compliance standards and company protocols",
        "Managed document organization, invoice reconciliation, and record keeping with meticulous attention to detail",
        "Maintained a clean, organized workspace and ensured all office areas met professional appearance standards",
        "Demonstrated reliability through consistent attendance and punctual reporting over multi-year tenure",
        "Developed strong client relationships through professional communication and follow-through",
    ]:
        add_bullet(doc, b)

    add_para(doc, "Sales Associate  |  Office Depot OfficeMax", bold=True, size=11, space_after=1)
    add_para(doc, "September 2020 \u2013 December 2020", size=10, space_after=4)
    for b in [
        "Delivered outstanding customer service by assisting customers with product selection and problem resolution",
        "Operated POS systems and handled cash transactions accurately and efficiently",
        "Maintained store appearance through floor cleaning, restocking, and organized product displays",
        "Applied upselling techniques to meet sales goals while prioritizing customer satisfaction",
    ]:
        add_bullet(doc, b)

    add_heading_styled(doc, "Education", level=2)
    add_para(doc, "Bachelor\u2019s Degree  |  University of Mississippi", bold=True, space_after=2)
    add_para(doc, "High School Diploma  |  Pulaski Academy", bold=True)

    add_heading_styled(doc, "Key Skills", level=2)
    skills = [
        "Vehicle & Facility Sanitation", "Workspace Maintenance & Organization",
        "Safety Inspection & Hazard Reporting", "Equipment Operation & Care",
        "Customer Service & Communication", "Inventory Management",
        "Team Collaboration & Reliability", "Heavy Lifting & Physical Tasks",
        "Attention to Detail", "Microsoft Office & Productivity Software",
        "POS Systems & Cash Handling",
    ]
    for i in range(0, len(skills), 2):
        left = skills[i]
        right = skills[i + 1] if i + 1 < len(skills) else ""
        text = f"\u2022  {left}"
        if right:
            text += f"     \u2022  {right}"
        add_para(doc, text, size=10, space_after=2)

    add_heading_styled(doc, "Certifications", level=2)
    add_para(doc, "Valid Driver\u2019s License (Clean Record)", bold=True)

    doc.save(os.path.join(OUT, "Chase_Kinslow_Resume.docx"))

# ─── COVER LETTER (Word) ────────────────────────────────────────

def make_cover_letter_docx():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    add_para(doc, "Chase Kinslow", bold=True, size=14, space_after=1)
    add_para(doc, "Monroe, LA", size=11, space_after=8)
    add_para(doc, "March 22, 2026", size=11, space_after=8)

    for line in ["Yolanda Johnson", "Talent Acquisition", "Enterprise Mobility", "1919 Cypress Street", "West Monroe, LA 71291"]:
        add_para(doc, line, size=11, space_after=1, bold=(line == "Yolanda Johnson"))
    doc.paragraphs[-1].paragraph_format.space_after = Pt(8)

    add_para(doc, "Dear Ms. Johnson,", size=11, space_after=8)

    paragraphs = [
        "I am writing to express my strong interest in the Automotive Detailer position at the Enterprise Mobility West Monroe location. With a background in operations, facility maintenance, and customer service, I am confident I can contribute to Enterprise Mobility\u2019s commitment to upholding the highest cleanliness and safety standards in the industry.",
        "In my current role in Operations & Logistics at Bayou DeSiard Country Club, I manage daily facility upkeep including cleaning, sanitation, and equipment maintenance to ensure a premium environment. This experience has sharpened my ability to maintain strict cleanliness protocols, inspect equipment for safety concerns, and keep organized, hazard-free workspaces\u2014all skills that directly align with the Automotive Detailer role. Prior to that, my years as an Operations & Compliance Associate at Little and Associates reinforced my attention to detail, reliability, and ability to follow established procedures with consistency.",
        "What draws me to Enterprise Mobility specifically is the company\u2019s founding values of integrity, honest work, and team spirit. I believe strongly in doing the right thing and showing up every day ready to give my best. I am also excited about Enterprise Mobility\u2019s promote-from-within culture and the opportunity to grow my career within an organization that invests in its people. Starting as an Automotive Detailer and contributing to the team\u2019s success while building toward future opportunities is exactly the kind of career path I am looking for.",
        "I meet all position qualifications: I am over 18, hold a valid driver\u2019s license with a clean driving record, am authorized to work in the United States, and have well over six consecutive months of work experience. I am available for the posted schedule, including rotating Saturdays, and am ready to start immediately.",
        "Thank you for considering my application. I would welcome the opportunity to discuss how my experience and work ethic can contribute to your West Monroe team.",
    ]
    for p in paragraphs:
        add_para(doc, p, size=11, space_after=8)

    add_para(doc, "Sincerely,", size=11, space_after=1)
    add_para(doc, "Chase Kinslow", bold=True, size=11)

    doc.save(os.path.join(OUT, "Chase_Kinslow_Cover_Letter.docx"))

# ─── EMAIL (Word) ────────────────────────────────────────────────

def make_email_docx():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    add_para(doc, "To: yolanda.johnson@ehi.com", bold=True, size=11, space_after=1)
    add_para(doc, "Subject: Automotive Detailer Application \u2013 West Monroe Location | Chase Kinslow", bold=True, size=11, space_after=12)

    add_para(doc, "Dear Ms. Johnson,", size=11, space_after=8)

    paragraphs = [
        "I hope this message finds you well. My name is Chase Kinslow and I am reaching out regarding the Automotive Detailer position at the West Monroe, LA location on Cypress Street.",
        "As a fellow Louisiana professional, I have a great deal of respect for the work you and your talent acquisition team do at Enterprise Mobility. I understand that under your leadership, Enterprise has earned the Candidate Experience Award multiple times, which speaks volumes about the kind of people-first culture you help build. That commitment to treating candidates with care and respect is exactly why I am drawn to this organization. It is clear that the values you champion in recruiting mirror Enterprise Mobility\u2019s founding principles of integrity, honest work, and team spirit.",
        "I have attached my resume and cover letter for your review. I bring hands-on experience in operations, facility maintenance, and sanitation from my current role at Bayou DeSiard Country Club, along with several years in operations and compliance work. I meet all of the posted qualifications and am available to start immediately on the posted schedule.",
        "Your background in connecting people with the right opportunities, and your dedication to mentoring careers, makes me confident that you understand the value of a candidate who is eager to grow from within. Starting as an Automotive Detailer and building a career at Enterprise Mobility is exactly the kind of path I am looking for, and I would be grateful for the chance to demonstrate my work ethic and commitment to your West Monroe team.",
        "I would love the opportunity to speak with you at your convenience. Thank you for your time and consideration.",
    ]
    for p in paragraphs:
        add_para(doc, p, size=11, space_after=8)

    add_para(doc, "Best regards,", size=11, space_after=1)
    add_para(doc, "Chase Kinslow", bold=True, size=11)

    doc.save(os.path.join(OUT, "Chase_Kinslow_Email.docx"))

# ─── PDF versions (reportlab) ───────────────────────────────────

DARK = HexColor("#1A1A2E")
BODY_COLOR = HexColor("#1E1E1E")

def get_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='Name', fontName='Helvetica-Bold', fontSize=20, textColor=DARK, alignment=1, spaceAfter=2))
    styles.add(ParagraphStyle(name='Subtitle', fontName='Helvetica', fontSize=10, textColor=HexColor("#505050"), alignment=1, spaceAfter=8))
    styles.add(ParagraphStyle(name='SectionHead', fontName='Helvetica-Bold', fontSize=13, textColor=DARK, spaceAfter=4, spaceBefore=10))
    styles.add(ParagraphStyle(name='Body', fontName='Helvetica', fontSize=10, textColor=BODY_COLOR, leading=14, spaceAfter=4))
    styles.add(ParagraphStyle(name='BodyBold', fontName='Helvetica-Bold', fontSize=10, textColor=BODY_COLOR, leading=14, spaceAfter=2))
    styles.add(ParagraphStyle(name='DateItalic', fontName='Helvetica-Oblique', fontSize=9, textColor=HexColor("#555555"), spaceAfter=4))
    styles.add(ParagraphStyle(name='BulletItem', fontName='Helvetica', fontSize=10, textColor=BODY_COLOR, leading=13, spaceAfter=2, leftIndent=18, bulletIndent=6))
    styles.add(ParagraphStyle(name='Letter', fontName='Helvetica', fontSize=11, textColor=BODY_COLOR, leading=15, spaceAfter=8))
    styles.add(ParagraphStyle(name='LetterBold', fontName='Helvetica-Bold', fontSize=11, textColor=BODY_COLOR, leading=15, spaceAfter=2))
    styles.add(ParagraphStyle(name='LetterName', fontName='Helvetica-Bold', fontSize=14, textColor=BODY_COLOR, spaceAfter=2))
    styles.add(ParagraphStyle(name='SkillItem', fontName='Helvetica', fontSize=10, textColor=BODY_COLOR, leading=13, spaceAfter=2))
    return styles

def section_hr():
    return HRFlowable(width="100%", thickness=1, color=DARK, spaceAfter=6)

def make_resume_pdf():
    s = get_styles()
    doc = SimpleDocTemplate(os.path.join(OUT, "Chase_Kinslow_Resume.pdf"), pagesize=letter,
                            topMargin=0.5*inch, bottomMargin=0.5*inch, leftMargin=0.65*inch, rightMargin=0.65*inch)
    story = []

    story.append(Paragraph("CHASE KINSLOW", s['Name']))
    story.append(Paragraph("Monroe, LA  |  Valid Driver's License", s['Subtitle']))
    story.append(Spacer(1, 6))

    story.append(Paragraph("PROFESSIONAL SUMMARY", s['SectionHead']))
    story.append(section_hr())
    story.append(Paragraph("Dedicated operations professional with hands-on experience in facility maintenance, sanitation, logistics, and customer service. Proven ability to maintain clean, organized workspaces while upholding safety and compliance standards. Strong attention to detail with a track record of reliability, punctuality, and team collaboration across fast-paced environments.", s['Body']))
    story.append(Spacer(1, 4))

    story.append(Paragraph("WORK EXPERIENCE", s['SectionHead']))
    story.append(section_hr())

    # Job 1
    story.append(Paragraph("Operations &amp; Logistics  |  Bayou DeSiard Country Club", s['BodyBold']))
    story.append(Paragraph("May 2025 - Present", s['DateItalic']))
    for b in [
        "Managed daily facility operations including cleaning, sanitation, and maintenance of indoor and outdoor spaces to uphold premium presentation standards",
        "Performed equipment inspections and ensured all tools and machinery were maintained in safe, working condition",
        "Coordinated logistics for events and daily operations, demonstrating strong organizational and multitasking skills",
        "Maintained clean and orderly work areas, identifying and reporting unsafe or hazardous conditions",
        "Collaborated with team members to deliver exceptional service to club members and guests",
    ]:
        story.append(Paragraph("\u2022  " + b, s['BulletItem']))
    story.append(Spacer(1, 4))

    # Job 2
    story.append(Paragraph("Operations &amp; Compliance Associate  |  Little and Associates", s['BodyBold']))
    story.append(Paragraph("December 2020 - 2025", s['DateItalic']))
    for b in [
        "Executed daily operational tasks with strict adherence to regulatory compliance standards and company protocols",
        "Managed document organization, invoice reconciliation, and record keeping with meticulous attention to detail",
        "Maintained a clean, organized workspace and ensured all office areas met professional appearance standards",
        "Demonstrated reliability through consistent attendance and punctual reporting over multi-year tenure",
        "Developed strong client relationships through professional communication and follow-through",
    ]:
        story.append(Paragraph("\u2022  " + b, s['BulletItem']))
    story.append(Spacer(1, 4))

    # Job 3
    story.append(Paragraph("Sales Associate  |  Office Depot OfficeMax", s['BodyBold']))
    story.append(Paragraph("September 2020 - December 2020", s['DateItalic']))
    for b in [
        "Delivered outstanding customer service by assisting customers with product selection and problem resolution",
        "Operated POS systems and handled cash transactions accurately and efficiently",
        "Maintained store appearance through floor cleaning, restocking, and organized product displays",
        "Applied upselling techniques to meet sales goals while prioritizing customer satisfaction",
    ]:
        story.append(Paragraph("\u2022  " + b, s['BulletItem']))
    story.append(Spacer(1, 4))

    # Education
    story.append(Paragraph("EDUCATION", s['SectionHead']))
    story.append(section_hr())
    story.append(Paragraph("Bachelor's Degree  |  University of Mississippi", s['BodyBold']))
    story.append(Paragraph("High School Diploma  |  Pulaski Academy", s['BodyBold']))
    story.append(Spacer(1, 4))

    # Skills
    story.append(Paragraph("KEY SKILLS", s['SectionHead']))
    story.append(section_hr())
    skills = [
        "Vehicle &amp; Facility Sanitation", "Workspace Maintenance &amp; Organization",
        "Safety Inspection &amp; Hazard Reporting", "Equipment Operation &amp; Care",
        "Customer Service &amp; Communication", "Inventory Management",
        "Team Collaboration &amp; Reliability", "Heavy Lifting &amp; Physical Tasks",
        "Attention to Detail", "Microsoft Office &amp; Productivity Software",
        "POS Systems &amp; Cash Handling",
    ]
    for i in range(0, len(skills), 2):
        left = skills[i]
        right = skills[i + 1] if i + 1 < len(skills) else ""
        txt = f"\u2022  {left}"
        if right:
            txt += f"&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;\u2022  {right}"
        story.append(Paragraph(txt, s['SkillItem']))
    story.append(Spacer(1, 4))

    # Certifications
    story.append(Paragraph("CERTIFICATIONS", s['SectionHead']))
    story.append(section_hr())
    story.append(Paragraph("Valid Driver's License (Clean Record)", s['BodyBold']))

    doc.build(story)


def make_letter_pdf(filename, elements_fn):
    s = get_styles()
    doc = SimpleDocTemplate(os.path.join(OUT, filename), pagesize=letter,
                            topMargin=0.75*inch, bottomMargin=0.75*inch, leftMargin=1*inch, rightMargin=1*inch)
    story = elements_fn(s)
    doc.build(story)


def make_cover_letter_pdf():
    def build(s):
        story = []
        story.append(Paragraph("Chase Kinslow", s['LetterName']))
        story.append(Paragraph("Monroe, LA", s['Letter']))
        story.append(Spacer(1, 4))
        story.append(Paragraph("March 22, 2026", s['Letter']))
        story.append(Spacer(1, 4))
        for line in ["Yolanda Johnson", "Talent Acquisition", "Enterprise Mobility", "1919 Cypress Street", "West Monroe, LA 71291"]:
            st = s['LetterBold'] if line == "Yolanda Johnson" else s['Letter']
            p = Paragraph(line, st)
            p.spaceAfter = 2
            story.append(p)
        story.append(Spacer(1, 6))
        story.append(Paragraph("Dear Ms. Johnson,", s['Letter']))
        story.append(Spacer(1, 2))

        paragraphs = [
            "I am writing to express my strong interest in the Automotive Detailer position at the Enterprise Mobility West Monroe location. With a background in operations, facility maintenance, and customer service, I am confident I can contribute to Enterprise Mobility's commitment to upholding the highest cleanliness and safety standards in the industry.",
            "In my current role in Operations &amp; Logistics at Bayou DeSiard Country Club, I manage daily facility upkeep including cleaning, sanitation, and equipment maintenance to ensure a premium environment. This experience has sharpened my ability to maintain strict cleanliness protocols, inspect equipment for safety concerns, and keep organized, hazard-free workspaces\u2014all skills that directly align with the Automotive Detailer role. Prior to that, my years as an Operations &amp; Compliance Associate at Little and Associates reinforced my attention to detail, reliability, and ability to follow established procedures with consistency.",
            "What draws me to Enterprise Mobility specifically is the company's founding values of integrity, honest work, and team spirit. I believe strongly in doing the right thing and showing up every day ready to give my best. I am also excited about Enterprise Mobility's promote-from-within culture and the opportunity to grow my career within an organization that invests in its people. Starting as an Automotive Detailer and contributing to the team's success while building toward future opportunities is exactly the kind of career path I am looking for.",
            "I meet all position qualifications: I am over 18, hold a valid driver's license with a clean driving record, am authorized to work in the United States, and have well over six consecutive months of work experience. I am available for the posted schedule, including rotating Saturdays, and am ready to start immediately.",
            "Thank you for considering my application. I would welcome the opportunity to discuss how my experience and work ethic can contribute to your West Monroe team.",
        ]
        for p in paragraphs:
            story.append(Paragraph(p, s['Letter']))

        story.append(Spacer(1, 4))
        story.append(Paragraph("Sincerely,", s['Letter']))
        story.append(Paragraph("Chase Kinslow", s['LetterBold']))
        return story

    make_letter_pdf("Chase_Kinslow_Cover_Letter.pdf", build)


def make_email_pdf():
    def build(s):
        story = []
        story.append(Paragraph("<b>To:</b> yolanda.johnson@ehi.com", s['Letter']))
        story.append(Paragraph("<b>Subject:</b> Automotive Detailer Application \u2013 West Monroe Location | Chase Kinslow", s['Letter']))
        story.append(Spacer(1, 8))
        story.append(Paragraph("Dear Ms. Johnson,", s['Letter']))
        story.append(Spacer(1, 2))

        paragraphs = [
            "I hope this message finds you well. My name is Chase Kinslow and I am reaching out regarding the Automotive Detailer position at the West Monroe, LA location on Cypress Street.",
            "As a fellow Louisiana professional, I have a great deal of respect for the work you and your talent acquisition team do at Enterprise Mobility. I understand that under your leadership, Enterprise has earned the Candidate Experience Award multiple times, which speaks volumes about the kind of people-first culture you help build. That commitment to treating candidates with care and respect is exactly why I am drawn to this organization. It is clear that the values you champion in recruiting mirror Enterprise Mobility's founding principles of integrity, honest work, and team spirit.",
            "I have attached my resume and cover letter for your review. I bring hands-on experience in operations, facility maintenance, and sanitation from my current role at Bayou DeSiard Country Club, along with several years in operations and compliance work. I meet all of the posted qualifications and am available to start immediately on the posted schedule.",
            "Your background in connecting people with the right opportunities, and your dedication to mentoring careers, makes me confident that you understand the value of a candidate who is eager to grow from within. Starting as an Automotive Detailer and building a career at Enterprise Mobility is exactly the kind of path I am looking for, and I would be grateful for the chance to demonstrate my work ethic and commitment to your West Monroe team.",
            "I would love the opportunity to speak with you at your convenience. Thank you for your time and consideration.",
        ]
        for p in paragraphs:
            story.append(Paragraph(p, s['Letter']))

        story.append(Spacer(1, 4))
        story.append(Paragraph("Best regards,", s['Letter']))
        story.append(Paragraph("Chase Kinslow", s['LetterBold']))
        return story

    make_letter_pdf("Chase_Kinslow_Email.pdf", build)


if __name__ == "__main__":
    make_resume_docx()
    make_cover_letter_docx()
    make_email_docx()
    make_resume_pdf()
    make_cover_letter_pdf()
    make_email_pdf()
    print("All 6 files generated successfully!")
    for f in ["Chase_Kinslow_Resume.docx", "Chase_Kinslow_Resume.pdf",
              "Chase_Kinslow_Cover_Letter.docx", "Chase_Kinslow_Cover_Letter.pdf",
              "Chase_Kinslow_Email.docx", "Chase_Kinslow_Email.pdf"]:
        path = os.path.join(OUT, f)
        size = os.path.getsize(path)
        print(f"  {f}: {size:,} bytes")
