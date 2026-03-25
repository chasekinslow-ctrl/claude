"""
Generate all application materials for JPMorgan Chase Transactions Specialist II (Monroe, LA):
1. Tailored Resume (PDF + DOCX)
2. Cover Letter (PDF + DOCX)
3. Cold Email to Steve Roop (TXT)
"""
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.colors import HexColor, white
from reportlab.pdfgen import canvas
from reportlab.lib.utils import simpleSplit
import os, textwrap

OUT = "/home/user/claude"

# ── Design tokens ──
NAVY = HexColor("#2B3544")
ORANGE = HexColor("#E8872B")
DARK_TEXT = HexColor("#1A1A1A")
WHITE = white
SIDEBAR_WIDTH = 2.1 * inch
PAGE_W, PAGE_H = letter

# ═══════════════════════════════════════════════════════════════
# PDF HELPER FUNCTIONS (shared by resume)
# ═══════════════════════════════════════════════════════════════

def draw_wrapped(c, text, x, y, max_width, font_name, font_size, leading, color=DARK_TEXT):
    c.setFont(font_name, font_size)
    c.setFillColor(color)
    lines = simpleSplit(text, font_name, font_size, max_width)
    for line in lines:
        if y < 40:
            c.showPage(); y = PAGE_H - 40; draw_sidebar_bg(c)
        c.setFont(font_name, font_size); c.setFillColor(color)
        c.drawString(x, y, line); y -= leading
    return y

def draw_bullet(c, text, x, y, max_width, indent=12):
    c.setFillColor(ORANGE); c.setFont("Helvetica", 10)
    c.drawString(x, y, "\u2014")
    c.setFillColor(DARK_TEXT); c.setFont("Helvetica", 9.5)
    text_x = x + indent
    lines = simpleSplit(text, "Helvetica", 9.5, max_width - indent)
    for i, line in enumerate(lines):
        if y < 40:
            c.showPage(); y = PAGE_H - 40; draw_sidebar_bg(c)
        c.setFont("Helvetica", 9.5); c.setFillColor(DARK_TEXT)
        c.drawString(text_x if i == 0 else text_x + 4, y, line); y -= 13
    return y

def draw_sidebar_bg(c):
    c.setFillColor(NAVY); c.rect(0, 0, SIDEBAR_WIDTH, PAGE_H, fill=1, stroke=0)

def draw_sidebar_section(c, title, y):
    c.setFillColor(ORANGE); c.setFont("Helvetica-Bold", 9.5)
    c.drawString(18, y, title); return y - 16

def draw_sidebar_item(c, text, y, bold=False):
    c.setFillColor(WHITE)
    font = "Helvetica-Bold" if bold else "Helvetica"
    c.setFont(font, 8.5)
    for line in simpleSplit(text, font, 8.5, SIDEBAR_WIDTH - 36):
        c.drawString(22, y, line); y -= 12
    return y

def draw_sidebar_bullet(c, text, y):
    c.setFillColor(WHITE); c.setFont("Helvetica", 8.5)
    lines = simpleSplit(text, "Helvetica", 8.5, SIDEBAR_WIDTH - 44)
    for i, line in enumerate(lines):
        c.drawString(22 if i == 0 else 32, y, ("\u2022  " + line) if i == 0 else line); y -= 12
    return y


# ═══════════════════════════════════════════════════════════════
# 1. RESUME — tailored for Transactions Specialist II
# ═══════════════════════════════════════════════════════════════

def make_resume_pdf():
    filepath = os.path.join(OUT, "Chase_Kinslow_Resume_JPMorganChase.pdf")
    c = canvas.Canvas(filepath, pagesize=letter)
    draw_sidebar_bg(c)

    # ── SIDEBAR ──
    sy = PAGE_H - 50

    sy = draw_sidebar_section(c, "CONTACT", sy)
    sy = draw_sidebar_item(c, "(501) 707-7779", sy)
    sy = draw_sidebar_item(c, "chasekn30@yahoo.com", sy)
    sy = draw_sidebar_item(c, "Monroe, LA 71201", sy)
    sy -= 10

    sy = draw_sidebar_section(c, "CORE COMPETENCIES", sy)
    for comp in [
        "Transaction Processing",
        "Document Handling & Data Entry",
        "Cash Handling & Reconciliation",
        "POS & Payment Systems",
        "Compliance & Regulatory Adherence",
        "Quality Control & Accuracy",
        "Inventory & Records Management",
        "Customer Service & Inquiries",
        "Problem Solving & Escalation",
        "Process Improvement",
        "Cross-Functional Teamwork",
    ]:
        sy = draw_sidebar_bullet(c, comp, sy)
    sy -= 10

    sy = draw_sidebar_section(c, "EDUCATION", sy)
    sy = draw_sidebar_item(c, "Bachelor of Accountancy", sy, bold=True)
    sy = draw_sidebar_item(c, "University of Mississippi", sy)
    sy -= 4
    sy = draw_sidebar_item(c, "High School Diploma", sy, bold=True)
    sy = draw_sidebar_item(c, "Pulaski Academy", sy)
    sy -= 10

    sy = draw_sidebar_section(c, "TOOLS & SYSTEMS", sy)
    for tool in [
        "Microsoft Office Suite",
        "QuickBooks Desktop & Online",
        "NCR POS Systems",
        "Digital Inventory Tracking",
        "Advanced Excel",
        "ADP Workforce Now",
    ]:
        sy = draw_sidebar_bullet(c, tool, sy)

    # ── MAIN CONTENT ──
    mx = SIDEBAR_WIDTH + 24
    mw = PAGE_W - mx - 30
    my = PAGE_H - 48

    c.setFont("Helvetica-Bold", 26); c.setFillColor(DARK_TEXT)
    c.drawString(mx, my, "CHASE KINSLOW"); my -= 22

    c.setFont("Helvetica-Oblique", 11); c.setFillColor(ORANGE)
    c.drawString(mx, my, "Transaction Processing & Operations Professional"); my -= 28

    # Professional Summary
    c.setFont("Helvetica-Bold", 11.5); c.setFillColor(DARK_TEXT)
    c.drawString(mx, my, "PROFESSIONAL SUMMARY"); my -= 4
    c.setStrokeColor(ORANGE); c.setLineWidth(1.5)
    c.line(mx, my, mx + mw, my); my -= 14

    summary = (
        "Detail-oriented operations professional with a Bachelor of Accountancy and hands-on "
        "experience in high-volume transaction processing, document handling, cash management, "
        "and regulatory compliance. Background spans zero-error reconciliations, digital tracking "
        "systems, data entry, and cross-departmental coordination in fast-paced environments. "
        "Proven ability to maintain production standards while ensuring accuracy, escalate "
        "non-routine issues effectively, and contribute to process improvement initiatives."
    )
    my = draw_wrapped(c, summary, mx, my, mw, "Helvetica", 9.5, 13, DARK_TEXT); my -= 10

    # Key Achievements & Skills
    c.setFont("Helvetica-Bold", 11.5); c.setFillColor(DARK_TEXT)
    c.drawString(mx, my, "KEY ACHIEVEMENTS & SKILLS"); my -= 4
    c.setStrokeColor(ORANGE); c.line(mx, my, mx + mw, my); my -= 16

    # Transaction Processing & Accuracy
    c.setFont("Helvetica-Bold", 10); c.setFillColor(DARK_TEXT)
    c.drawString(mx, my, "Transaction Processing, Data Entry & Accuracy"); my -= 14
    for b in [
        "Executed high-volume POS transactions with zero-error accuracy, managing daily cash handling, register reconciliation, and payment processing across multiple product categories",
        "Managed comprehensive documentation workflows requiring zero-error precision \u2014 reconciliations, workpapers, and submissions where a single discrepancy disqualified an entire project file",
        "Directed inventory control using digital tracking systems to prevent shrink, maintain accurate stock levels, and ensure supplies were distributed across all operational areas",
    ]:
        my = draw_bullet(c, b, mx, my, mw); my -= 2
    my -= 6

    # Customer Service & Escalation
    c.setFont("Helvetica-Bold", 10); c.setFillColor(DARK_TEXT)
    c.drawString(mx, my, "Customer Service, Inquiries & Escalation"); my -= 14
    for b in [
        "Handled customer inquiries and walk-in consultations, matching needs to available solutions while delivering a positive experience at every touchpoint",
        "Cross-trained across all departments to provide coverage wherever needed, ensuring uninterrupted service during peak periods and staff absences",
        "Escalated non-routine issues to senior team members, applying common sense and experience from similar situations to identify potential solutions",
    ]:
        my = draw_bullet(c, b, mx, my, mw); my -= 2
    my -= 6

    # Compliance & Process Improvement
    c.setFont("Helvetica-Bold", 10); c.setFillColor(DARK_TEXT)
    c.drawString(mx, my, "Compliance, Reconciliation & Process Improvement"); my -= 14
    for b in [
        "Orchestrated monthly close procedures with rigorous reconciliation of accounts and records to maintain total data integrity across multiple client portfolios",
        "Performed systematic opening and closing facility audits, applying inspection protocols to ensure compliance with brand standards, safety requirements, and operational benchmarks",
        "Maintained up-to-date knowledge of regulatory requirements and internal procedures, applying this knowledge to ensure compliance in all daily activities",
    ]:
        my = draw_bullet(c, b, mx, my, mw); my -= 2
    my -= 10

    # Employment History
    c.setFont("Helvetica-Bold", 11.5); c.setFillColor(DARK_TEXT)
    c.drawString(mx, my, "EMPLOYMENT HISTORY"); my -= 4
    c.setStrokeColor(ORANGE); c.line(mx, my, mx + mw, my); my -= 16

    for title, company, dates in [
        ("Operations & Logistics", "Bayou DeSiard Country Club", "2025 \u2013 Present"),
        ("Accountant", "Little and Associates, LLC", "2020 \u2013 2023"),
        ("Sales Associate", "Office Depot", "2020"),
    ]:
        c.setFont("Helvetica-Bold", 9.5); c.setFillColor(DARK_TEXT)
        lt = f"{title}  |  "; c.drawString(mx, my, lt)
        tw = c.stringWidth(lt, "Helvetica-Bold", 9.5)
        c.setFont("Helvetica-Oblique", 9.5); c.drawString(mx + tw, my, company)
        tw2 = c.stringWidth(company, "Helvetica-Oblique", 9.5)
        c.setFont("Helvetica", 9.5); c.drawString(mx + tw + tw2, my, f"  |  {dates}")
        my -= 16

    c.save()
    print(f"Resume PDF: {filepath} ({os.path.getsize(filepath):,} bytes)")


def make_resume_docx():
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'; style.font.size = Pt(10)
    for s in doc.sections:
        s.top_margin = Inches(0.5); s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.6); s.right_margin = Inches(0.6)

    O = RGBColor(0xE8, 0x87, 0x2B); D = RGBColor(0x1A, 0x1A, 0x1A)

    def orange_line(doc):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(4)
        pPr = p._element.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr')
        b = OxmlElement('w:bottom'); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '6')
        b.set(qn('w:space'), '1'); b.set(qn('w:color'), 'E8872B')
        pBdr.append(b); pPr.append(pBdr)

    def sec(doc, t):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t); r.font.size = Pt(12); r.bold = True; r.font.color.rgb = D; orange_line(doc)

    def subsec(doc, t):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(t); r.font.size = Pt(10.5); r.bold = True; r.font.color.rgb = D

    def bul(doc, t):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.left_indent = Pt(18)
        r = p.add_run("\u2014  " + t); r.font.size = Pt(9.5); r.font.color.rgb = D

    # Name
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after = Pt(0)
    r = p.add_run("CHASE KINSLOW"); r.font.size = Pt(24); r.bold = True; r.font.color.rgb = D

    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Transaction Processing & Operations Professional"); r.font.size = Pt(11); r.italic = True; r.font.color.rgb = O

    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
    r = p.add_run("(501) 707-7779  |  chasekn30@yahoo.com  |  Monroe, LA 71201"); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    sec(doc, "PROFESSIONAL SUMMARY")
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(
        "Detail-oriented operations professional with a Bachelor of Accountancy and hands-on "
        "experience in high-volume transaction processing, document handling, cash management, "
        "and regulatory compliance. Background spans zero-error reconciliations, digital tracking "
        "systems, data entry, and cross-departmental coordination in fast-paced environments. "
        "Proven ability to maintain production standards while ensuring accuracy, escalate "
        "non-routine issues effectively, and contribute to process improvement initiatives."
    ); r.font.size = Pt(9.5)

    sec(doc, "KEY ACHIEVEMENTS & SKILLS")

    subsec(doc, "Transaction Processing, Data Entry & Accuracy")
    for b in [
        "Executed high-volume POS transactions with zero-error accuracy, managing daily cash handling, register reconciliation, and payment processing across multiple product categories",
        "Managed comprehensive documentation workflows requiring zero-error precision \u2014 reconciliations, workpapers, and submissions where a single discrepancy disqualified an entire project file",
        "Directed inventory control using digital tracking systems to prevent shrink, maintain accurate stock levels, and ensure supplies were distributed across all operational areas",
    ]: bul(doc, b)

    subsec(doc, "Customer Service, Inquiries & Escalation")
    for b in [
        "Handled customer inquiries and walk-in consultations, matching needs to available solutions while delivering a positive experience at every touchpoint",
        "Cross-trained across all departments to provide coverage wherever needed, ensuring uninterrupted service during peak periods and staff absences",
        "Escalated non-routine issues to senior team members, applying common sense and experience from similar situations to identify potential solutions",
    ]: bul(doc, b)

    subsec(doc, "Compliance, Reconciliation & Process Improvement")
    for b in [
        "Orchestrated monthly close procedures with rigorous reconciliation of accounts and records to maintain total data integrity across multiple client portfolios",
        "Performed systematic opening and closing facility audits, applying inspection protocols to ensure compliance with brand standards, safety requirements, and operational benchmarks",
        "Maintained up-to-date knowledge of regulatory requirements and internal procedures, applying this knowledge to ensure compliance in all daily activities",
    ]: bul(doc, b)

    sec(doc, "EMPLOYMENT HISTORY")
    for title, company, dates in [
        ("Operations & Logistics", "Bayou DeSiard Country Club", "2025 \u2013 Present"),
        ("Accountant", "Little and Associates, LLC", "2020 \u2013 2023"),
        ("Sales Associate", "Office Depot", "2020"),
    ]:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = D
        r = p.add_run("  |  "); r.font.size = Pt(9.5)
        r = p.add_run(company); r.italic = True; r.font.size = Pt(9.5)
        r = p.add_run(f"  |  {dates}"); r.font.size = Pt(9.5)

    sec(doc, "CORE COMPETENCIES")
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run("  \u2022  ".join([
        "Transaction Processing", "Document Handling & Data Entry",
        "Cash Handling & Reconciliation", "POS & Payment Systems",
        "Compliance & Regulatory Adherence", "Quality Control & Accuracy",
        "Inventory & Records Management", "Customer Service & Inquiries",
        "Problem Solving & Escalation", "Process Improvement",
        "Cross-Functional Teamwork",
    ])); r.font.size = Pt(9)

    fp = os.path.join(OUT, "Chase_Kinslow_Resume_JPMorganChase.docx")
    doc.save(fp)
    print(f"Resume DOCX: {fp} ({os.path.getsize(fp):,} bytes)")


# ═══════════════════════════════════════════════════════════════
# 2. COVER LETTER
# ═══════════════════════════════════════════════════════════════

COVER_LETTER = """Chase Kinslow
Monroe, LA 71201
(501) 707-7779 | chasekn30@yahoo.com

March 25, 2026

JPMorgan Chase & Co.
Hiring Manager \u2014 Transactions Specialist II
Monroe, LA

Dear Hiring Manager,

I am writing to express my strong interest in the Transactions Specialist II position at JPMorgan Chase in Monroe, Louisiana. With a Bachelor of Accountancy from the University of Mississippi and direct experience in high-volume transaction processing, document handling, cash reconciliation, and regulatory compliance, I am confident I can contribute to your operations team from day one.

In my current role at Bayou DeSiard Country Club, I manage daily operations and logistics that require the same precision and structured execution this role demands \u2014 processing transactions accurately, maintaining production standards, and coordinating across departments to keep operations running smoothly. Previously, as an Accountant at Little and Associates, LLC, I managed documentation workflows where zero-error precision was non-negotiable: reconciliations, workpapers, and client submissions where a single discrepancy could disqualify an entire project file. That experience built the exact discipline required for processing and clearing transactions with the accuracy JPMorgan Chase expects.

What draws me to this opportunity specifically is Chase\u2019s 30-year investment in North Louisiana \u2014 with nearly 1,000 employees at the Monroe campus and the recent opening of the Ruston Operations Center, it\u2019s clear the firm is committed to growing operations here. I want to be part of that growth. I am also excited by JPMorgan Chase\u2019s push toward AI and automation in transaction processing. I am eager to develop proficiency with these technologies and contribute to innovation efforts that optimize how work gets done.

I bring strong critical thinking and problem-solving skills, a track record of handling customer inquiries with a positive-experience mindset, and the ability to escalate non-routine issues by applying common sense and experience from similar situations. I am comfortable sitting for extended periods, maintaining focus on detail-intensive work, and operating in a structured, supervised environment.

I would welcome the opportunity to discuss how my background in transaction processing, compliance, and operations aligns with what you need on your team. Thank you for your consideration.

Sincerely,
Chase Kinslow"""


def make_cover_letter_pdf():
    filepath = os.path.join(OUT, "Chase_Kinslow_CoverLetter_JPMorganChase.pdf")
    c = canvas.Canvas(filepath, pagesize=letter)
    x, y = 72, PAGE_H - 72
    mw = PAGE_W - 144

    def write_line(text, font="Helvetica", size=10.5, spacing=14, color=DARK_TEXT):
        nonlocal y
        c.setFont(font, size); c.setFillColor(color)
        if not text.strip():
            y -= spacing; return
        lines = simpleSplit(text, font, size, mw)
        for line in lines:
            c.drawString(x, y, line); y -= spacing

    # Header
    write_line("Chase Kinslow", "Helvetica-Bold", 14, 16)
    write_line("Monroe, LA 71201", "Helvetica", 9.5, 13, HexColor("#555555"))
    write_line("(501) 707-7779  |  chasekn30@yahoo.com", "Helvetica", 9.5, 13, HexColor("#555555"))
    y -= 16

    write_line("March 25, 2026")
    y -= 8
    write_line("JPMorgan Chase & Co.")
    write_line("Hiring Manager \u2014 Transactions Specialist II")
    write_line("Monroe, LA")
    y -= 8
    write_line("Dear Hiring Manager,")
    y -= 4

    paragraphs = [
        "I am writing to express my strong interest in the Transactions Specialist II position at JPMorgan Chase in Monroe, Louisiana. With a Bachelor of Accountancy from the University of Mississippi and direct experience in high-volume transaction processing, document handling, cash reconciliation, and regulatory compliance, I am confident I can contribute to your operations team from day one.",

        "In my current role at Bayou DeSiard Country Club, I manage daily operations and logistics that require the same precision and structured execution this role demands \u2014 processing transactions accurately, maintaining production standards, and coordinating across departments to keep operations running smoothly. Previously, as an Accountant at Little and Associates, LLC, I managed documentation workflows where zero-error precision was non-negotiable: reconciliations, workpapers, and client submissions where a single discrepancy could disqualify an entire project file. That experience built the exact discipline required for processing and clearing transactions with the accuracy JPMorgan Chase expects.",

        "What draws me to this opportunity specifically is Chase\u2019s 30-year investment in North Louisiana \u2014 with nearly 1,000 employees at the Monroe campus and the recent opening of the Ruston Operations Center, it\u2019s clear the firm is committed to growing operations here. I want to be part of that growth. I am also excited by JPMorgan Chase\u2019s push toward AI and automation in transaction processing. I am eager to develop proficiency with these technologies and contribute to innovation efforts that optimize how work gets done.",

        "I bring strong critical thinking and problem-solving skills, a track record of handling customer inquiries with a positive-experience mindset, and the ability to escalate non-routine issues by applying common sense and experience from similar situations. I am comfortable sitting for extended periods, maintaining focus on detail-intensive work, and operating in a structured, supervised environment.",

        "I would welcome the opportunity to discuss how my background in transaction processing, compliance, and operations aligns with what you need on your team. Thank you for your consideration.",
    ]

    for para in paragraphs:
        write_line(para)
        y -= 6

    y -= 4
    write_line("Sincerely,")
    write_line("Chase Kinslow", "Helvetica-Bold", 10.5)

    c.save()
    print(f"Cover Letter PDF: {filepath} ({os.path.getsize(filepath):,} bytes)")


def make_cover_letter_docx():
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor

    doc = Document()
    style = doc.styles['Normal']; style.font.name = 'Calibri'; style.font.size = Pt(10.5)
    for s in doc.sections:
        s.top_margin = Inches(1); s.bottom_margin = Inches(1)
        s.left_margin = Inches(1); s.right_margin = Inches(1)

    for line in COVER_LETTER.strip().split('\n'):
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(2)
        if line.startswith("Chase Kinslow") and "Monroe" not in line and "|" not in line:
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(14)

    fp = os.path.join(OUT, "Chase_Kinslow_CoverLetter_JPMorganChase.docx")
    doc.save(fp)
    print(f"Cover Letter DOCX: {fp} ({os.path.getsize(fp):,} bytes)")


# ═══════════════════════════════════════════════════════════════
# 3. COLD EMAIL TO STEVE ROOP
# ═══════════════════════════════════════════════════════════════

EMAIL_TEXT = """Subject: Congratulations on Business Leader of the Year \u2014 and a Question About Your Monroe Team

Mr. Roop,

First, congratulations on being named the Greater Shreveport Chamber\u2019s 2025 Business Leader of the Year. That recognition \u2014 after 40+ years of leadership in financial services and the kind of civic investment you\u2019ve made across North Louisiana, from the Ark-La-Tex Air Service Alliance to your Honorary Commander role with the 307th Bomb Wing \u2014 is well deserved.

I\u2019m reaching out because I live in Monroe and recently applied for the Transactions Specialist II position at the Monroe Operations Center. I wanted to introduce myself directly.

I have a Bachelor of Accountancy from Ole Miss and have spent the last several years in operations and compliance roles here in the Monroe area \u2014 managing zero-error reconciliations, high-volume transaction processing, inventory control with digital tracking systems, and cross-departmental coordination. My current role at Bayou DeSiard Country Club has me running daily operations and logistics, and before that I spent three years at an accounting firm where a single discrepancy in a workpaper could disqualify an entire client file. That kind of precision is exactly what your transaction processing team needs.

What excites me most about this opportunity is Chase\u2019s long-term commitment to North Louisiana. Nearly 1,000 employees at the Monroe campus, the new $31 million Ruston Operations Center now open and growing toward 200 jobs \u2014 this isn\u2019t a company passing through. I want to build a career here, contribute to the team\u2019s production and accuracy standards, and grow with the operation as JPMorgan Chase continues to invest in this region.

I\u2019ve attached my resume for your reference. If there\u2019s someone specific on the Monroe operations team you\u2019d recommend I connect with, I\u2019d be grateful for the introduction. Either way, thank you for your time and for everything you do for this community.

Respectfully,
Chase Kinslow
(501) 707-7779
chasekn30@yahoo.com
Monroe, LA 71201"""


def write_email():
    fp = os.path.join(OUT, "Cold_Email_Steve_Roop.txt")
    with open(fp, 'w') as f:
        f.write(EMAIL_TEXT.strip())
    print(f"Email: {fp} ({os.path.getsize(fp):,} bytes)")


# ═══════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════

if __name__ == "__main__":
    make_resume_pdf()
    make_resume_docx()
    make_cover_letter_pdf()
    make_cover_letter_docx()
    write_email()
    print("\nAll JPMorgan Chase application materials generated!")
