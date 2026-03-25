"""
Regenerate JPMorgan Chase resume matching the gray-blue sidebar template format:
- Light steel-blue sidebar (not dark navy)
- Gray background subsection header bars
- » arrow bullets in sidebar
- • bullet points in main content
- Orange/gold accent underlines on main section headers
- CREDENTIALS and ADDITIONAL sections in sidebar
"""
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.colors import HexColor, white, black
from reportlab.pdfgen import canvas
from reportlab.lib.utils import simpleSplit
import os

OUT = "/home/user/claude"

# ── Colors matching the gray-blue template ──
SIDEBAR_BG = HexColor("#8899A6")  # steel blue-gray
SIDEBAR_HEADER = HexColor("#2A2A2A")  # dark headers in sidebar
GOLD_LINE = HexColor("#C8A951")  # gold/amber accent line
GRAY_BAR = HexColor("#D8D8D8")  # gray subsection background bar
DARK_TEXT = HexColor("#1A1A1A")
LIGHT_TEXT = HexColor("#FFFFFF")
SIDEBAR_TEXT = HexColor("#F0F0F0")  # slightly off-white for sidebar body text
PAGE_W, PAGE_H = letter
SIDEBAR_WIDTH = 2.0 * inch


def draw_sidebar_bg(c):
    c.setFillColor(SIDEBAR_BG)
    c.rect(0, 0, SIDEBAR_WIDTH, PAGE_H, fill=1, stroke=0)


def draw_sidebar_section(c, title, y):
    """Draw a sidebar section header with letter spacing."""
    c.setFillColor(SIDEBAR_HEADER)
    c.setFont("Helvetica-Bold", 8)
    # Add letter spacing by drawing spaced chars
    spaced = "   ".join(list(title))
    c.drawString(14, y, spaced)
    return y - 14


def draw_sidebar_item(c, text, y, bold=False):
    c.setFillColor(SIDEBAR_TEXT)
    font = "Helvetica-Bold" if bold else "Helvetica"
    c.setFont(font, 7.5)
    lines = simpleSplit(text, font, 7.5, SIDEBAR_WIDTH - 28)
    for line in lines:
        c.drawString(16, y, line)
        y -= 11
    return y


def draw_sidebar_arrow(c, text, y):
    """Draw a » arrow bullet in the sidebar."""
    c.setFillColor(SIDEBAR_TEXT)
    c.setFont("Helvetica", 7.5)
    lines = simpleSplit(text, "Helvetica", 7.5, SIDEBAR_WIDTH - 36)
    for i, line in enumerate(lines):
        if i == 0:
            c.drawString(16, y, "\u00BB" + " " + line)
        else:
            c.drawString(26, y, line)
        y -= 11
    return y


def draw_main_section(c, title, y, mx, mw):
    """Draw a main section header with gold underline."""
    c.setFont("Helvetica-Bold", 11)
    c.setFillColor(DARK_TEXT)
    c.drawString(mx, y, title)
    y -= 4
    c.setStrokeColor(GOLD_LINE)
    c.setLineWidth(1.5)
    c.line(mx, y, mx + mw, y)
    return y - 12


def draw_subsection_bar(c, title, y, mx, mw):
    """Draw a gray background bar with subsection title."""
    bar_h = 14
    c.setFillColor(GRAY_BAR)
    c.rect(mx - 4, y - 3, mw + 8, bar_h, fill=1, stroke=0)
    c.setFont("Helvetica-Bold", 9)
    c.setFillColor(DARK_TEXT)
    c.drawString(mx, y, title)
    return y - 16


def draw_bullet(c, text, x, y, max_width):
    """Draw a • bullet point with wrapped text."""
    c.setFillColor(DARK_TEXT)
    c.setFont("Helvetica", 8.5)
    c.drawString(x, y, "\u2022")
    text_x = x + 10
    lines = simpleSplit(text, "Helvetica", 8.5, max_width - 10)
    for i, line in enumerate(lines):
        if y < 40:
            c.showPage()
            y = PAGE_H - 40
            draw_sidebar_bg(c)
        c.setFont("Helvetica", 8.5)
        c.setFillColor(DARK_TEXT)
        c.drawString(text_x if i == 0 else text_x + 4, y, line)
        y -= 12
    return y


def draw_wrapped(c, text, x, y, max_width, font="Helvetica", size=8.5, leading=12, color=DARK_TEXT):
    c.setFont(font, size)
    c.setFillColor(color)
    lines = simpleSplit(text, font, size, max_width)
    for line in lines:
        if y < 40:
            c.showPage(); y = PAGE_H - 40; draw_sidebar_bg(c)
        c.setFont(font, size); c.setFillColor(color)
        c.drawString(x, y, line)
        y -= leading
    return y


def make_resume():
    filepath = os.path.join(OUT, "Chase_Kinslow_Resume_JPMorganChase.pdf")
    c = canvas.Canvas(filepath, pagesize=letter)

    draw_sidebar_bg(c)

    # ══════════ SIDEBAR ══════════
    sy = PAGE_H - 40

    sy = draw_sidebar_section(c, "CONTACT", sy)
    sy = draw_sidebar_item(c, "(501) 707-7779", sy)
    sy = draw_sidebar_item(c, "chasekn30@yahoo.com", sy)
    sy = draw_sidebar_item(c, "Monroe, LA 71201", sy)
    sy -= 10

    sy = draw_sidebar_section(c, "CORE COMPETENCIES", sy)
    for comp in [
        "Transaction Processing",
        "Document Handling & Data Entry",
        "Cash Handling & Reconciliation",
        "POS & Payment Systems",
        "Compliance Verification",
        "Quality Control & Accuracy",
        "Inventory & Records Management",
        "Customer Service & Inquiries",
        "Audit-Ready Record Keeping",
    ]:
        sy = draw_sidebar_arrow(c, comp, sy)
    sy -= 10

    sy = draw_sidebar_section(c, "CREDENTIALS", sy)
    sy = draw_sidebar_item(c, "Licensed CPA", sy, bold=True)
    sy = draw_sidebar_item(c, "Arkansas & Louisiana", sy)
    sy -= 10

    sy = draw_sidebar_section(c, "EDUCATION", sy)
    sy = draw_sidebar_item(c, "Bachelor of Accountancy", sy, bold=True)
    sy = draw_sidebar_item(c, "University of Mississippi", sy)
    sy -= 10

    sy = draw_sidebar_section(c, "ADDITIONAL", sy)
    for item in [
        "Valid Driver's License \u2014 Clean MVR",
        "Reliable Personal Transportation",
        "Available for Extended Hours",
        "Advanced Excel Proficiency",
        "Microsoft Office Suite",
        "QuickBooks Desktop & Online",
        "NCR POS Systems",
    ]:
        sy = draw_sidebar_arrow(c, item, sy)

    # ══════════ MAIN CONTENT ══════════
    mx = SIDEBAR_WIDTH + 20
    mw = PAGE_W - mx - 28
    my = PAGE_H - 38

    # Name
    c.setFont("Helvetica-Bold", 24)
    c.setFillColor(DARK_TEXT)
    c.drawString(mx, my, "CHASE KINSLOW")
    my -= 16

    # Subtitle
    c.setFont("Helvetica", 9)
    c.setFillColor(HexColor("#555555"))
    c.drawString(mx, my, "Transaction Processing & Operations Professional")
    my -= 20

    # Professional Summary
    my = draw_main_section(c, "PROFESSIONAL SUMMARY", my, mx, mw)

    summary = (
        "Detail-oriented operations professional with a Bachelor of Accountancy and hands-on "
        "experience in high-volume transaction processing, document handling, cash management, "
        "and regulatory compliance. Proven ability to verify work against contracts, maintain "
        "audit-ready records under deadline pressure, and flag discrepancies before they escalate. "
        "Background combines rigorous financial documentation discipline with operations and "
        "logistics coordination. Local to Monroe, LA with reliable transportation and immediate "
        "availability for extended-hour, 7-day activations."
    )
    my = draw_wrapped(c, summary, mx, my, mw)
    my -= 8

    # Key Achievements & Skills
    my = draw_main_section(c, "KEY ACHIEVEMENTS & SKILLS", my, mx, mw)

    # Sub: Transaction Processing & Accuracy
    my = draw_subsection_bar(c, "Transaction Processing, Data Entry & Accuracy", my, mx, mw)
    for b in [
        "Executed high-volume POS transactions with zero-error accuracy, managing daily cash handling, register reconciliation, and payment processing across multiple product categories",
        "Managed comprehensive documentation workflows requiring zero-error precision \u2014 reconciliations, workpapers, and submissions where a single discrepancy disqualified an entire project file",
        "Directed inventory control using digital tracking systems to prevent shrink, maintain accurate stock levels, and ensure supplies were distributed across all operational areas",
    ]:
        my = draw_bullet(c, b, mx, my, mw)
        my -= 2
    my -= 4

    # Sub: Customer Service & Escalation
    my = draw_subsection_bar(c, "Customer Service, Inquiries & Escalation", my, mx, mw)
    for b in [
        "Handled customer inquiries and walk-in consultations, matching needs to available solutions while delivering a positive experience at every touchpoint",
        "Cross-trained across all departments to provide coverage wherever needed, ensuring uninterrupted service during peak periods and staff absences",
        "Escalated non-routine issues to senior team members, applying common sense and experience from similar situations to identify potential solutions",
    ]:
        my = draw_bullet(c, b, mx, my, mw)
        my -= 2
    my -= 4

    # Sub: Compliance & Process Improvement
    my = draw_subsection_bar(c, "Compliance, Reconciliation & Process Improvement", my, mx, mw)
    for b in [
        "Orchestrated monthly close procedures with rigorous reconciliation of accounts and records to maintain total data integrity across multiple client portfolios",
        "Performed systematic opening and closing facility audits, applying inspection protocols to ensure compliance with brand standards, safety requirements, and operational benchmarks",
        "Maintained up-to-date knowledge of regulatory requirements and internal procedures, applying this knowledge to ensure compliance in all daily activities",
    ]:
        my = draw_bullet(c, b, mx, my, mw)
        my -= 2
    my -= 8

    # Employment History
    my = draw_main_section(c, "EMPLOYMENT HISTORY", my, mx, mw)

    for title, company, dates in [
        ("Operations & Logistics", "Bayou DeSiard Country Club", "2025"),
        ("Accountant", "Little and Associates, LLC", "2020 \u2013 2023"),
        ("Sales Associate", "Office Depot", "2020"),
    ]:
        c.setFont("Helvetica-Bold", 8.5)
        c.setFillColor(DARK_TEXT)
        c.drawString(mx, my, title)
        tw = c.stringWidth(title, "Helvetica-Bold", 8.5)
        c.setFont("Helvetica-Oblique", 8.5)
        c.setFillColor(HexColor("#555555"))
        c.drawString(mx + tw, my, f" | {company}")
        # Right-align dates
        c.setFont("Helvetica", 8.5)
        dw = c.stringWidth(dates, "Helvetica", 8.5)
        c.drawString(mx + mw - dw, my, dates)
        my -= 14

    c.save()
    print(f"Resume PDF: {filepath} ({os.path.getsize(filepath):,} bytes)")


def make_resume_docx():
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    style = doc.styles['Normal']; style.font.name = 'Calibri'; style.font.size = Pt(10)
    for s in doc.sections:
        s.top_margin = Inches(0.5); s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.6); s.right_margin = Inches(0.6)

    GOLD = RGBColor(0xC8, 0xA9, 0x51); D = RGBColor(0x1A, 0x1A, 0x1A); G = RGBColor(0x55, 0x55, 0x55)

    def gold_line(doc):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(4)
        pPr = p._element.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr')
        b = OxmlElement('w:bottom'); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '6')
        b.set(qn('w:space'), '1'); b.set(qn('w:color'), 'C8A951')
        pBdr.append(b); pPr.append(pBdr)

    def sec(doc, t):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t); r.font.size = Pt(11); r.bold = True; r.font.color.rgb = D; gold_line(doc)

    def subsec(doc, t):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
        # Gray background via shading
        pPr = p._element.get_or_add_pPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'D8D8D8')
        pPr.append(shd)
        r = p.add_run(t); r.font.size = Pt(9.5); r.bold = True; r.font.color.rgb = D

    def bul(doc, t):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.left_indent = Pt(18)
        r = p.add_run("\u2022  " + t); r.font.size = Pt(9); r.font.color.rgb = D

    # Name
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after = Pt(0)
    r = p.add_run("CHASE KINSLOW"); r.font.size = Pt(22); r.bold = True; r.font.color.rgb = D

    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Transaction Processing & Operations Professional"); r.font.size = Pt(9); r.font.color.rgb = G

    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
    r = p.add_run("(501) 707-7779  |  chasekn30@yahoo.com  |  Monroe, LA 71201"); r.font.size = Pt(8.5); r.font.color.rgb = G

    sec(doc, "PROFESSIONAL SUMMARY")
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(
        "Detail-oriented operations professional with a Bachelor of Accountancy and hands-on "
        "experience in high-volume transaction processing, document handling, cash management, "
        "and regulatory compliance. Proven ability to verify work against contracts, maintain "
        "audit-ready records under deadline pressure, and flag discrepancies before they escalate. "
        "Background combines rigorous financial documentation discipline with operations and "
        "logistics coordination. Local to Monroe, LA with reliable transportation and immediate "
        "availability for extended-hour, 7-day activations."
    ); r.font.size = Pt(9)

    sec(doc, "KEY ACHIEVEMENTS & SKILLS")

    subsec(doc, "Transaction Processing, Data Entry & Accuracy")
    for b in [
        "Executed high-volume POS transactions with zero-error accuracy, managing daily cash handling, register reconciliation, and payment processing across multiple product categories",
        "Managed comprehensive documentation workflows requiring zero-error precision \u2014 reconciliations, workpapers, and submissions where a single discrepancy disqualified an entire project file",
        "Directed inventory control using digital tracking systems to prevent shrink, maintain accurate stock levels, and ensure supplies were distributed across all operational areas",
    ]: bul(doc, b)

    subsec(doc, "Customer Service, Inquiries & Escalation")
    for b in [
        "Handled customer inquiries and walk-in consultations, matching needs to available solutions while delivering a positive experience at every touchpoint",
        "Cross-trained across all departments to provide coverage wherever needed, ensuring uninterrupted service during peak periods and staff absences",
        "Escalated non-routine issues to senior team members, applying common sense and experience from similar situations to identify potential solutions",
    ]: bul(doc, b)

    subsec(doc, "Compliance, Reconciliation & Process Improvement")
    for b in [
        "Orchestrated monthly close procedures with rigorous reconciliation of accounts and records to maintain total data integrity across multiple client portfolios",
        "Performed systematic opening and closing facility audits, applying inspection protocols to ensure compliance with brand standards, safety requirements, and operational benchmarks",
        "Maintained up-to-date knowledge of regulatory requirements and internal procedures, applying this knowledge to ensure compliance in all daily activities",
    ]: bul(doc, b)

    sec(doc, "EMPLOYMENT HISTORY")
    for title, company, dates in [
        ("Operations & Logistics", "Bayou DeSiard Country Club", "2025"),
        ("Accountant", "Little and Associates, LLC", "2020 \u2013 2023"),
        ("Sales Associate", "Office Depot", "2020"),
    ]:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = D
        r = p.add_run(f" | {company}"); r.font.size = Pt(9); r.font.color.rgb = G
        r = p.add_run(f"    {dates}"); r.font.size = Pt(9); r.font.color.rgb = G

    sec(doc, "CORE COMPETENCIES")
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run("  \u2022  ".join([
        "Transaction Processing", "Document Handling & Data Entry",
        "Cash Handling & Reconciliation", "POS & Payment Systems",
        "Compliance Verification", "Quality Control & Accuracy",
        "Inventory & Records Management", "Customer Service & Inquiries",
        "Audit-Ready Record Keeping",
    ])); r.font.size = Pt(8.5)

    fp = os.path.join(OUT, "Chase_Kinslow_Resume_JPMorganChase.docx")
    doc.save(fp)
    print(f"Resume DOCX: {fp} ({os.path.getsize(fp):,} bytes)")


if __name__ == "__main__":
    make_resume()
    make_resume_docx()
    print("Done!")
